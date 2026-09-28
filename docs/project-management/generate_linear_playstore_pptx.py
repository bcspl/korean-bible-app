#!/usr/bin/env python3
"""Linear-style Play Store release guide PPTX for Korean Bible App."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt
from lxml import etree

# --- Linear design tokens ---
BG = RGBColor(0x08, 0x09, 0x0A)
SURFACE = RGBColor(0x11, 0x12, 0x14)
SURFACE2 = RGBColor(0x16, 0x17, 0x1A)
BORDER = RGBColor(0x27, 0x29, 0x2D)
ACCENT = RGBColor(0x5E, 0x6A, 0xD2)
ACCENT_SOFT = RGBColor(0x8B, 0x93, 0xE8)
TEXT = RGBColor(0xF7, 0xF8, 0xF8)
MUTED = RGBColor(0x8A, 0x8F, 0x98)
DIM = RGBColor(0x5C, 0x61, 0x6B)
GREEN = RGBColor(0x4C, 0xB7, 0x82)
AMBER = RGBColor(0xF2, 0xC9, 0x4C)
RED = RGBColor(0xEB, 0x57, 0x57)

W = Inches(13.333)
H = Inches(7.5)
MARGIN = Inches(0.55)
TODAY = date.today().isoformat()


def _set_run(run, size=14, bold=False, color=TEXT, font="Segoe UI"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_rect(slide, x, y, w, h, fill=SURFACE, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        x, y, w, h,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def add_text(slide, x, y, w, h, text, size=14, bold=False, color=TEXT, align=PP_ALIGN.LEFT, font="Segoe UI", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(2 if font == "Consolas" else 0)
        run = p.add_run()
        run.text = part
        _set_run(run, size=size, bold=bold, color=color, font=font)
    try:
        bodyPr = tf._txBody.bodyPr
        if valign == MSO_ANCHOR.MIDDLE:
            bodyPr.set("anchor", "ctr")
        elif valign == MSO_ANCHOR.BOTTOM:
            bodyPr.set("anchor", "b")
        else:
            bodyPr.set("anchor", "t")
    except Exception:
        pass
    return box


def add_lines(slide, x, y, w, h, lines, size=13, color=MUTED, bold_first=False, line_spacing=1.15):
    """lines: list of str or (str, dict)"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(opts.get("before", 2))
        p.space_after = Pt(opts.get("after", 4))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size=opts.get("size", size),
            bold=opts.get("bold", bold_first and i == 0),
            color=opts.get("color", color),
            font=opts.get("font", "Segoe UI"),
        )
    return box


def footer(slide, page, total):
    add_text(slide, MARGIN, H - Inches(0.38), Inches(8), Inches(0.28),
             "Korean Bible App  ·  Play Store Release Guide  ·  Linear style",
             size=10, color=DIM)
    add_text(slide, W - MARGIN - Inches(1.2), H - Inches(0.38), Inches(1.2), Inches(0.28),
             f"{page} / {total}", size=10, color=DIM, align=PP_ALIGN.RIGHT)


def header(slide, eyebrow, title, subtitle=None):
    add_text(slide, MARGIN, Inches(0.28), Inches(10), Inches(0.28),
             eyebrow.upper(), size=11, bold=True, color=ACCENT_SOFT)
    add_text(slide, MARGIN, Inches(0.52), Inches(12), Inches(0.5),
             title, size=28, bold=True, color=TEXT)
    if subtitle:
        add_text(slide, MARGIN, Inches(1.05), Inches(12), Inches(0.35),
                 subtitle, size=13, color=MUTED)


def pill(slide, x, y, w, h, text, fill=ACCENT, text_color=TEXT):
    add_rect(slide, x, y, w, h, fill=fill, radius=True)
    add_text(slide, x, y, w, h, text, size=11, bold=True, color=text_color,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def status_dot_label(slide, x, y, label, color):
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y + Inches(0.06), Inches(0.14), Inches(0.14))
    oval.fill.solid()
    oval.fill.fore_color.rgb = color
    oval.line.fill.background()
    add_text(slide, x + Inches(0.22), y, Inches(1.6), Inches(0.28), label, size=11, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    slides_meta = []

    def new():
        s = prs.slides.add_slide(blank)
        add_bg(s)
        slides_meta.append(s)
        return s

    # ========== 1 Title ==========
    s = new()
    # left accent bar
    add_rect(s, 0, 0, Inches(0.12), H, fill=ACCENT)
    add_text(s, MARGIN, Inches(1.8), Inches(11), Inches(0.35),
             "PROJECT PLAYBOOK", size=12, bold=True, color=ACCENT_SOFT)
    add_text(s, MARGIN, Inches(2.2), Inches(12), Inches(1.0),
             "앞으로 할 작업 & Play Store 출시 가이드", size=36, bold=True, color=TEXT)
    add_text(s, MARGIN, Inches(3.3), Inches(11), Inches(0.5),
             "한국어 성경 앱  ·  Public Domain  ·  목표 공개 2026-11", size=16, color=MUTED)
    # cards
    cards = [
        ("제품", "~92%", "기능·콘텐츠"),
        ("스토어 준비", "~20%", "서명·정책·리스팅"),
        ("다음 차단", "3건", "ID · Key · Privacy"),
    ]
    for i, (a, b, c) in enumerate(cards):
        x = MARGIN + i * Inches(3.9)
        add_rect(s, x, Inches(4.3), Inches(3.6), Inches(1.55), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, x + Inches(0.25), Inches(4.45), Inches(3.1), Inches(0.3), a, size=12, color=MUTED)
        add_text(s, x + Inches(0.25), Inches(4.8), Inches(3.1), Inches(0.5), b, size=28, bold=True, color=TEXT)
        add_text(s, x + Inches(0.25), Inches(5.4), Inches(3.1), Inches(0.3), c, size=12, color=DIM)
    add_text(s, MARGIN, Inches(6.2), Inches(10), Inches(0.3),
             f"Updated {TODAY}  ·  Linear-inspired dark UI", size=11, color=DIM)

    # ========== 2 Agenda ==========
    s = new()
    header(s, "Contents", "이 문서에서 다루는 것", "작업 백로그와 Play 출시 절차를 한 덱에 정리")
    items = [
        ("01", "현재 상태와 차단 항목"),
        ("02", "앞으로 할 작업 (Sprint A–D)"),
        ("03", "Play Store end-to-end 절차 (Phase 0–8)"),
        ("04", "로컬 서명 · AAB 빌드 명령"),
        ("05", "스토어 리스팅 · 정책 · 심사"),
        ("06", "월별 로드맵 · 즉시 Next Actions"),
    ]
    for i, (num, title) in enumerate(items):
        row, col = divmod(i, 2)
        x = MARGIN + col * Inches(6.2)
        y = Inches(1.7) + row * Inches(1.35)
        add_rect(s, x, y, Inches(5.9), Inches(1.15), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, x + Inches(0.3), y + Inches(0.28), Inches(1.0), Inches(0.5), num, size=22, bold=True, color=ACCENT)
        add_text(s, x + Inches(1.3), y + Inches(0.35), Inches(4.3), Inches(0.5), title, size=16, bold=True, color=TEXT)

    # ========== 3 Status ==========
    s = new()
    header(s, "Status", "현재 위치", "제품은 거의 완성 · 공개 준비가 병목")
    left = [
        ("✅ 완료", GREEN, [
            "KRV · KJV · ASV 병렬 읽기",
            "PD 찬송 102 · 교독 51 · 기도/신경",
            "본문 검증 · flutter test PASS",
            "다크모드 · 글자크기 · 고대비",
            "gdrive_upload MCP (백업)",
        ]),
        ("⏳ 미완", AMBER, [
            "packageId (com.example.* 금지)",
            "키스토어 · AAB · 실기기 서명 테스트",
            "개인정보처리방침 URL",
            "온보딩 · 역본 저장 · 절 복사",
            "스크린샷 · 스토어 카피 · Console",
        ]),
    ]
    for i, (title, color, bullets) in enumerate(left):
        x = MARGIN + i * Inches(6.2)
        add_rect(s, x, Inches(1.65), Inches(5.9), Inches(4.7), fill=SURFACE, line=BORDER, radius=True)
        add_rect(s, x, Inches(1.65), Inches(0.12), Inches(4.7), fill=color)
        add_text(s, x + Inches(0.35), Inches(1.85), Inches(5.2), Inches(0.4), title, size=18, bold=True, color=TEXT)
        lines = [(f"  {b}", {"size": 14, "color": MUTED, "after": 10}) for b in bullets]
        add_lines(s, x + Inches(0.35), Inches(2.4), Inches(5.2), Inches(3.6), lines)

    # ========== 4 Sprint overview ==========
    s = new()
    header(s, "Backlog", "앞으로 할 작업 — Sprint 개요", "가치 ÷ 노력 · 11월 공개에 맞춘 순서")
    sprints = [
        ("A", "스토어 기술", "9월", ACCENT, ["packageId", "키스토어", "AAB", "실기기", "targetSdk"]),
        ("B", "공개 전 UX", "10월", ACCENT_SOFT, ["온보딩", "역본 저장", "복사/공유", "라이선스", "스모크"]),
        ("C", "자산·정책", "9–10월", AMBER, ["Privacy URL", "설명 카피", "아이콘/피처", "스크린샷", "Console"]),
        ("D", "공개 후", "12월+", DIM, ["통독", "위치복원", "하이라이트", "TTS", "2열"]),
    ]
    for i, (code, name, when, color, items) in enumerate(sprints):
        x = MARGIN + i * Inches(3.1)
        add_rect(s, x, Inches(1.65), Inches(2.95), Inches(4.7), fill=SURFACE, line=BORDER, radius=True)
        pill(s, x + Inches(0.2), Inches(1.85), Inches(0.55), Inches(0.32), code, fill=color, text_color=BG if color == AMBER else TEXT)
        add_text(s, x + Inches(0.85), Inches(1.85), Inches(1.9), Inches(0.32), when, size=12, color=MUTED, valign=MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(0.2), Inches(2.35), Inches(2.55), Inches(0.4), name, size=16, bold=True, color=TEXT)
        lines = [(f"·  {t}", {"size": 13, "color": MUTED, "after": 8}) for t in items]
        add_lines(s, x + Inches(0.2), Inches(2.9), Inches(2.55), Inches(3.0), lines)

    # ========== 5 Journey ==========
    s = new()
    header(s, "Play Store", "출시 여정 한눈에", "Phase 0 → 8 · 막히면 이전 Phase부터")
    phases = [
        ("0", "결정"),
        ("1", "Console"),
        ("2", "서명"),
        ("3", "빌드"),
        ("4", "리스팅"),
        ("5", "정책"),
        ("6", "테스트"),
        ("7", "심사"),
        ("8", "운영"),
    ]
    y = Inches(2.0)
    total_w = Inches(12.2)
    cell = total_w / len(phases)
    for i, (n, label) in enumerate(phases):
        x = MARGIN + i * cell
        add_rect(s, x, y, cell - Inches(0.12), Inches(1.35), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, x, y + Inches(0.25), cell - Inches(0.12), Inches(0.4), n, size=20, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        add_text(s, x, y + Inches(0.75), cell - Inches(0.12), Inches(0.4), label, size=12, color=TEXT, align=PP_ALIGN.CENTER)
        if i < len(phases) - 1:
            add_text(s, x + cell - Inches(0.18), y + Inches(0.5), Inches(0.2), Inches(0.3), "›", size=16, color=DIM)
    notes = [
        ("차단", "ID · Keystore · Privacy URL이 없으면 Phase 2/5에서 정지", RED),
        ("권장", "내부 테스트 트랙으로 먼저 설치 검증 후 프로덕션", GREEN),
        ("절대", "키스토어·key.properties 커밋 금지 · com.example.* 업로드 금지", AMBER),
    ]
    for i, (t, d, c) in enumerate(notes):
        y2 = Inches(3.7) + i * Inches(0.85)
        add_rect(s, MARGIN, y2, Inches(12.2), Inches(0.72), fill=SURFACE, line=BORDER, radius=True)
        add_rect(s, MARGIN, y2, Inches(0.1), Inches(0.72), fill=c)
        add_text(s, MARGIN + Inches(0.35), y2 + Inches(0.2), Inches(1.2), Inches(0.35), t, size=14, bold=True, color=c)
        add_text(s, MARGIN + Inches(1.7), y2 + Inches(0.2), Inches(10), Inches(0.35), d, size=14, color=MUTED)

    # ========== 6 Phase 0-1 ==========
    s = new()
    header(s, "Phase 0–1", "사전 결정 · Play Console 계정", "출시 전 반드시 확정할 입력값")
    add_rect(s, MARGIN, Inches(1.6), Inches(6.0), Inches(4.9), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(0.3), Inches(1.8), Inches(5.4), Inches(0.4), "Phase 0 — 결정 체크리스트", size=16, bold=True, color=TEXT)
    decisions = [
        "Application ID (평생 고정, reverse DNS)",
        "표시명: 한국어 성경",
        "지원 이메일",
        "Privacy Policy URL 호스팅처",
        "출시: Play only / APK 사이드로드",
        "키스토어 생성 경로 (예: C:\\Users\\LSH\\secure)",
    ]
    add_lines(s, MARGIN + Inches(0.3), Inches(2.4), Inches(5.4), Inches(3.8),
              [(f"☐  {d}", {"size": 13, "color": MUTED, "after": 10}) for d in decisions])

    add_rect(s, MARGIN + Inches(6.25), Inches(1.6), Inches(6.0), Inches(4.9), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(6.55), Inches(1.8), Inches(5.4), Inches(0.4), "Phase 1 — Console 절차", size=16, bold=True, color=TEXT)
    steps = [
        "1. play.google.com/console 가입 ($25)",
        "2. 앱 만들기 · 언어 한국어 · 무료",
        "3. 앱/게임 · 카테고리 초안 선택",
        "4. 정책 선언 설문 시작",
        "5. 내부 테스트 트랙 먼저 생성",
        "6. 테스터 이메일 추가 준비",
    ]
    add_lines(s, MARGIN + Inches(6.55), Inches(2.4), Inches(5.4), Inches(3.8),
              [(st, {"size": 13, "color": MUTED, "after": 10}) for st in steps])

    # ========== 7 Phase 2 Signing ==========
    s = new()
    header(s, "Phase 2", "Package ID · 업로드 키 · 서명", "Play가 거부하는 com.example.* 제거가 1순위")
    add_rect(s, MARGIN, Inches(1.6), Inches(12.2), Inches(1.1), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(0.3), Inches(1.75), Inches(11.5), Inches(0.35),
             "현재: com.example.korean_bible_app  →  목표 예: com.lsh.koreanbible", size=15, bold=True, color=TEXT)
    add_text(s, MARGIN + Inches(0.3), Inches(2.15), Inches(11.5), Inches(0.35),
             "android/app/build.gradle.kts 의 applicationId · namespace 변경 후 clean build", size=13, color=MUTED)

    cmds = [
        ("키스토어 생성 (1회)",
         "keytool -genkey -v\n"
         "-keystore C:\\Users\\LSH\\secure\\korean-bible-upload.jks\n"
         "-keyalg RSA -keysize 2048 -validity 10000\n"
         "-alias upload"),
        ("key.properties (gitignore)",
         "storePassword=***\n"
         "keyPassword=***\n"
         "keyAlias=upload\n"
         "storeFile=C:/Users/LSH/secure/korean-bible-upload.jks"),
    ]
    for i, (title, body) in enumerate(cmds):
        x = MARGIN + i * Inches(6.2)
        add_rect(s, x, Inches(2.95), Inches(5.95), Inches(3.35), fill=SURFACE2, line=BORDER, radius=True)
        add_text(s, x + Inches(0.25), Inches(3.1), Inches(5.4), Inches(0.35), title, size=14, bold=True, color=ACCENT_SOFT)
        add_text(s, x + Inches(0.25), Inches(3.55), Inches(5.4), Inches(2.5), body, size=12, color=MUTED, font="Consolas")

    # ========== 8 Phase 3 Build ==========
    s = new()
    header(s, "Phase 3", "품질 게이트 · AAB / APK 빌드", "콘텐츠 변경 이후 release 재빌드 필수")
    gate = [
        ("flutter test", "유닛/위젯 테스트"),
        ("verify_bible_texts.py", "성경 본문 무결성"),
        ("flutter analyze", "정적 분석 0 errors"),
        ("기기 스모크", "비행기모드 · 다크 · 큰글씨"),
    ]
    for i, (a, b) in enumerate(gate):
        x = MARGIN + (i % 4) * Inches(3.1)
        y = Inches(1.65)
        add_rect(s, x, y, Inches(2.95), Inches(1.35), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, x + Inches(0.2), y + Inches(0.3), Inches(2.55), Inches(0.4), a, size=13, bold=True, color=TEXT)
        add_text(s, x + Inches(0.2), y + Inches(0.75), Inches(2.55), Inches(0.35), b, size=12, color=MUTED)

    add_rect(s, MARGIN, Inches(3.25), Inches(12.2), Inches(3.15), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(0.3), Inches(3.45), Inches(11.5), Inches(0.35), "PowerShell 명령", size=14, bold=True, color=ACCENT_SOFT)
    build_txt = (
        "cd \"C:\\Users\\LSH\\Bible App\\korean-bible-app\"\n"
        "flutter test\n"
        "python scripts\\verify_bible_texts.py\n"
        "flutter build appbundle --release\n"
        "# -> build\\app\\outputs\\bundle\\release\\app-release.aab\n"
        "flutter build apk --release\n"
        "adb install -r build\\app\\outputs\\flutter-apk\\app-release.apk"
    )
    add_text(s, MARGIN + Inches(0.3), Inches(3.9), Inches(11.5), Inches(2.3), build_txt, size=13, color=MUTED, font="Consolas")

    # ========== 9 Phase 4-5 Listing ==========
    s = new()
    header(s, "Phase 4–5", "스토어 리스팅 · 정책 · 개인정보", "심사에서 가장 자주 막히는 구간")
    assets = [
        ("아이콘", "512×512 PNG"),
        ("피처 그래픽", "1024×500"),
        ("스크린샷", "폰 4–8장"),
        ("짧은 설명", "≤ 80자"),
        ("긴 설명", "≤ 4000자"),
        ("등급", "Everyone"),
    ]
    for i, (a, b) in enumerate(assets):
        x = MARGIN + (i % 6) * Inches(2.05)
        add_rect(s, x, Inches(1.6), Inches(1.95), Inches(1.15), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, x + Inches(0.1), Inches(1.75), Inches(1.75), Inches(0.35), a, size=12, bold=True, color=TEXT, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(2.2), Inches(1.75), Inches(0.35), b, size=11, color=MUTED, align=PP_ALIGN.CENTER)

    add_rect(s, MARGIN, Inches(3.0), Inches(6.0), Inches(3.4), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(0.25), Inches(3.2), Inches(5.5), Inches(0.35), "카피 초안", size=14, bold=True, color=TEXT)
    add_lines(s, MARGIN + Inches(0.25), Inches(3.7), Inches(5.5), Inches(2.5), [
        ("짧은", {"size": 12, "bold": True, "color": ACCENT_SOFT, "after": 4}),
        ("완전 오프라인 한국어 성경 · 찬송 · 예배자료", {"size": 13, "color": MUTED, "after": 12}),
        ("긴 설명 필수 문구", {"size": 12, "bold": True, "color": ACCENT_SOFT, "after": 4}),
        ("KRV·KJV·ASV · PD 찬송 · 광고/계정 없음", {"size": 13, "color": MUTED, "after": 4}),
        ("개역개정 미수록 · 한국찬송가 공식 아님", {"size": 13, "color": MUTED, "after": 4}),
    ])

    add_rect(s, MARGIN + Inches(6.25), Inches(3.0), Inches(6.0), Inches(3.4), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(6.5), Inches(3.2), Inches(5.5), Inches(0.35), "정책 체크", size=14, bold=True, color=TEXT)
    add_lines(s, MARGIN + Inches(6.5), Inches(3.7), Inches(5.5), Inches(2.5), [
        ("☐ Privacy Policy URL 공개", {"size": 13, "color": MUTED, "after": 8}),
        ("☐ 데이터 보안 양식 (수집 없음)", {"size": 13, "color": MUTED, "after": 8}),
        ("☐ 앱 콘텐츠 / 광고 없음 선언", {"size": 13, "color": MUTED, "after": 8}),
        ("☐ 민감 권한 최소화 확인", {"size": 13, "color": MUTED, "after": 8}),
        ("☐ 콘텐츠 고지 (PD · 비공식 찬송)", {"size": 13, "color": MUTED, "after": 8}),
    ])

    # ========== 10 Phase 6-7 ==========
    s = new()
    header(s, "Phase 6–7", "업로드 · 테스트 트랙 · 심사 · 공개", "한 번에 프로덕션보다 내부 테스트 권장")
    tracks = [
        ("1. 내부 테스트", "빠른 설치 검증 · 소수 이메일", ACCENT),
        ("2. 비공개/오픈", "피드백 · 크래시 관찰", ACCENT_SOFT),
        ("3. 프로덕션", "전체 사용자 공개", GREEN),
    ]
    for i, (t, d, c) in enumerate(tracks):
        x = MARGIN + i * Inches(4.1)
        add_rect(s, x, Inches(1.65), Inches(3.9), Inches(1.6), fill=SURFACE, line=BORDER, radius=True)
        add_rect(s, x, Inches(1.65), Inches(0.1), Inches(1.6), fill=c)
        add_text(s, x + Inches(0.3), Inches(1.9), Inches(3.4), Inches(0.4), t, size=15, bold=True, color=TEXT)
        add_text(s, x + Inches(0.3), Inches(2.45), Inches(3.4), Inches(0.5), d, size=13, color=MUTED)

    flow = [
        "AAB 업로드 → 버전 코드(+N) 증가 확인",
        "한국어 릴리스 노트 작성",
        "대시보드 ‘게시 전 필수 항목’ 전부 완료",
        "프로덕션 제출 → 심사 (신규 계정은 수일~수주)",
        "반려 시: 정책/권한/설명 수정 후 재제출",
        "공개 후: Vitals · 리뷰 · 핫픽스 AAB",
    ]
    add_rect(s, MARGIN, Inches(3.55), Inches(12.2), Inches(2.85), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(0.3), Inches(3.75), Inches(11.5), Inches(0.35), "제출 플로우", size=14, bold=True, color=TEXT)
    add_lines(s, MARGIN + Inches(0.3), Inches(4.25), Inches(11.5), Inches(2.0),
              [(f"{i+1}.  {t}", {"size": 14, "color": MUTED, "after": 6}) for i, t in enumerate(flow)])

    # ========== 11 Phase 8 + UX ==========
    s = new()
    header(s, "Phase 8 + UX", "버전 운영 · 공개 전 최소 UX", "1.0에 넣을 것 / 나중에 넣을 것")
    add_rect(s, MARGIN, Inches(1.6), Inches(6.0), Inches(4.9), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(0.3), Inches(1.8), Inches(5.4), Inches(0.4), "공개 전 최소 UX (Sprint B)", size=16, bold=True, color=TEXT)
    ux = [
        ("온보딩", "PD · 3역본 · 개역개정 미수록 3줄"),
        ("역본 저장", "SharedPreferences로 재선택 제거"),
        ("절 복사/공유", "소그룹 · 카톡 실용성"),
        ("라이선스 화면", "스토어 문의·신뢰"),
        ("스모크", "비행기모드 · 접근성"),
    ]
    for i, (a, b) in enumerate(ux):
        y = Inches(2.4) + i * Inches(0.7)
        add_text(s, MARGIN + Inches(0.35), y, Inches(1.8), Inches(0.35), a, size=13, bold=True, color=ACCENT_SOFT)
        add_text(s, MARGIN + Inches(2.2), y, Inches(3.5), Inches(0.35), b, size=13, color=MUTED)

    add_rect(s, MARGIN + Inches(6.25), Inches(1.6), Inches(6.0), Inches(4.9), fill=SURFACE, line=BORDER, radius=True)
    add_text(s, MARGIN + Inches(6.55), Inches(1.8), Inches(5.4), Inches(0.4), "운영 · 보안", size=16, bold=True, color=TEXT)
    ops = [
        "pubspec version: 1.0.0+1 → 매 업로드 +N",
        "Play App Signing 사용 권장",
        "키스토어 오프라인 백업 + 암호 관리자",
        "분실 시 업데이트 매우 어려움",
        "gitignore: *.jks, key.properties",
        "공개 후: 통독·하이라이트·TTS는 선택",
    ]
    add_lines(s, MARGIN + Inches(6.55), Inches(2.4), Inches(5.4), Inches(3.8),
              [(f"·  {o}", {"size": 13, "color": MUTED, "after": 10}) for o in ops])

    # ========== 12 Timeline ==========
    s = new()
    header(s, "Roadmap", "월별 로드맵 → 2026-11", "M4 기술 → M5 품질 → M6 공개")
    months = [
        ("8월", "문서 · 재빌드 · 결정 수집 · Drive", "NOW", ACCENT),
        ("9월", "ID · Key · AAB · Privacy · Console", "M4", ACCENT_SOFT),
        ("10월", "UX · 스모크 · 스크린샷 · 내부테스트", "M5", AMBER),
        ("11월", "프로덕션 제출 · 심사 · 1.0 공개", "M6", GREEN),
    ]
    for i, (m, d, tag, c) in enumerate(months):
        y = Inches(1.7) + i * Inches(1.15)
        add_rect(s, MARGIN, y, Inches(12.2), Inches(1.0), fill=SURFACE, line=BORDER, radius=True)
        add_rect(s, MARGIN, y, Inches(0.12), Inches(1.0), fill=c)
        add_text(s, MARGIN + Inches(0.4), y + Inches(0.28), Inches(1.4), Inches(0.45), m, size=18, bold=True, color=TEXT)
        add_text(s, MARGIN + Inches(2.0), y + Inches(0.32), Inches(8.5), Inches(0.4), d, size=15, color=MUTED)
        pill(s, W - MARGIN - Inches(1.5), y + Inches(0.3), Inches(1.3), Inches(0.4), tag, fill=c, text_color=BG if c == AMBER else TEXT)

    # ========== 13 Next actions ==========
    s = new()
    header(s, "Next", "지금 바로 할 일", "의사결정만 끝나면 기술 작업 착수 가능")
    actions = [
        ("1", "packageId 문자열 확정", "예: com.lsh.koreanbible", RED),
        ("2", "키스토어 생성 실행", "secure 폴더 + 비밀번호 보관", AMBER),
        ("3", "Privacy URL 위치 결정", "GitHub Pages / Notion 등", AMBER),
        ("4", "지원 이메일 · Play 계정", "스토어 리스팅 필수", ACCENT),
        ("5", "release AAB 1회 빌드", "내부 테스트 업로드", ACCENT_SOFT),
        ("6", "온보딩 + 역본 저장", "공개 전 UX 최소셋", GREEN),
    ]
    for i, (n, t, d, c) in enumerate(actions):
        col = i % 3
        row = i // 3
        x = MARGIN + col * Inches(4.1)
        y = Inches(1.65) + row * Inches(2.35)
        add_rect(s, x, y, Inches(3.9), Inches(2.1), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, x + Inches(0.25), y + Inches(0.3), Inches(0.5), Inches(0.45), n, size=24, bold=True, color=c)
        add_text(s, x + Inches(0.25), y + Inches(0.9), Inches(3.4), Inches(0.4), t, size=15, bold=True, color=TEXT)
        add_text(s, x + Inches(0.25), y + Inches(1.4), Inches(3.4), Inches(0.4), d, size=12, color=MUTED)

    # ========== 14 Appendix refs ==========
    s = new()
    header(s, "Appendix", "관련 문서 · 산출물", "MD가 원본 · PPTX는 공유용")
    refs = [
        ("PLAYSTORE_RELEASE_GUIDE_2026-08-10.md", "본 가이드 원문"),
        ("RELEASE_CHECKLIST_PLAYSTORE.md", "R1–R25 체크리스트"),
        ("DEPLOY_CHECKLIST.md", "빌드 명령 · 결정 질문"),
        ("VALUE_AND_USABILITY.md", "UX 백로그"),
        ("VISION_AND_PLAN.md", "비전 · 마일스톤"),
        ("APPLIED_AND_ROADMAP_PLAN_2026-08-10.md", "적용/계획 종합"),
    ]
    for i, (a, b) in enumerate(refs):
        y = Inches(1.65) + i * Inches(0.7)
        add_rect(s, MARGIN, y, Inches(12.2), Inches(0.58), fill=SURFACE, line=BORDER, radius=True)
        add_text(s, MARGIN + Inches(0.3), y + Inches(0.12), Inches(8), Inches(0.35), a, size=13, bold=True, color=TEXT, font="Consolas")
        add_text(s, MARGIN + Inches(8.3), y + Inches(0.12), Inches(3.6), Inches(0.35), b, size=13, color=MUTED)

    # ========== 15 Close ==========
    s = new()
    add_rect(s, 0, 0, Inches(0.12), H, fill=ACCENT)
    add_text(s, MARGIN, Inches(2.3), Inches(12), Inches(0.4), "READY WHEN YOU ARE", size=12, bold=True, color=ACCENT_SOFT)
    add_text(s, MARGIN, Inches(2.8), Inches(12), Inches(0.8),
             "packageId만 알려주시면\n서명·AAB 설정부터 바로 진행합니다.", size=28, bold=True, color=TEXT)
    add_text(s, MARGIN, Inches(4.5), Inches(12), Inches(0.4),
             "문서: docs/project-management/  ·  PPTX: Desktop/Bible Project/", size=14, color=MUTED)
    add_text(s, MARGIN, Inches(5.2), Inches(12), Inches(0.35),
             f"Korean Bible App  ·  {TODAY}", size=12, color=DIM)

    total = len(slides_meta)
    for i, slide in enumerate(slides_meta):
        if i == 0 or i == total - 1:
            continue
        footer(slide, i + 1, total)

    out_dirs = [
        Path.home() / "Desktop" / "Bible Project",
        Path(r"C:\Users\LSH\Bible App\korean-bible-app\docs\roadmap-pptx"),
        Path(r"C:\Users\LSH\Bible App\korean-bible-app\docs\project-management"),
    ]
    name = f"KoreanBible_PlayStore_ReleaseGuide_{TODAY}.pptx"
    paths = []
    for d in out_dirs:
        d.mkdir(parents=True, exist_ok=True)
        p = d / name
        prs.save(str(p))
        paths.append(p)
    return paths


if __name__ == "__main__":
    paths = build()
    for p in paths:
        print("WROTE", p)
