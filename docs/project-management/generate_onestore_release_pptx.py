#!/usr/bin/env python3
"""ONE store release process deck (Korean) for 한국어 성경 — 2026-10-03.
Style/helpers reused from generate_android_release_pptx.py (2026-09-29).

Usage:  python docs/project-management/generate_onestore_release_pptx.py [output.pptx]
Default output: docs/roadmap-pptx/KoreanBible_ONEstore_Release_Process_2026-10-03.pptx
Requires: python-pptx (pip install python-pptx)
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt, Emu

# ---------- design tokens ----------
FONT = "Malgun Gothic"
MONO = "Consolas"
BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LINE = RGBColor(0xE3, 0xE7, 0xEE)
CARD = RGBColor(0xF5, 0xF7, 0xFB)
ACCENT = RGBColor(0x3F, 0x51, 0xB5)      # app indigo
ACCENT_LT = RGBColor(0xE8, 0xEB, 0xF8)
CODE_BG = RGBColor(0x1E, 0x23, 0x2E)
CODE_FG = RGBColor(0xE6, 0xED, 0xF3)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
GREEN_LT = RGBColor(0xE6, 0xF4, 0xEA)
AMBER = RGBColor(0xB2, 0x6A, 0x00)
AMBER_LT = RGBColor(0xFE, 0xF3, 0xE0)
RED = RGBColor(0xC6, 0x28, 0x28)
RED_LT = RGBColor(0xFD, 0xEC, 0xEA)
BLUE = RGBColor(0x1A, 0x5F, 0xB4)
BLUE_LT = RGBColor(0xE7, 0xF0, 0xFB)

STATUS = {
    "DONE": ("완료", GREEN, GREEN_LT),
    "NEXT": ("다음", BLUE, BLUE_LT),
    "BLOCKED": ("사용자 필요", RED, RED_LT),
    "WAIT": ("대기", AMBER, AMBER_LT),
    "PART": ("진행 중", AMBER, AMBER_LT),
    "DRAFT": ("초안 완료", GREEN, GREEN_LT),
    "LOCK": ("변경 불가", RED, RED_LT),
    "CAUTION": ("주의", AMBER, AMBER_LT),
    "FREE": ("일반", GREEN, GREEN_LT),
}

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)
CONTENT_TOP = Inches(1.55)
MIN_PT = 14

DATE = "2026-10-03"
FOOTER = "한국어 성경 · 원스토어 출시 프로세스 · " + DATE

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
slides_meta: list = []
ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "store-assets" / "onestore"


def _font(run, size, bold=False, color=INK, font=FONT):
    assert size >= MIN_PT, size
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
        if el is None:
            el = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1], {})
            rPr.append(el)
        el.set("typeface", font)


def rect(s, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.RECTANGLE, radius=0.06):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, content, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.15, after=4, bullets=False):
    """content: str (\n = new paragraph) or list of str / (str, dict)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    items = content if isinstance(content, list) else str(content).split("\n")
    for i, it in enumerate(items):
        t, o = (it if isinstance(it, tuple) else (it, {}))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.line_spacing = o.get("spacing", spacing)
        p.space_after = Pt(o.get("after", after))
        prefix = "• " if (bullets and not o.get("nobullet")) else ""
        r = p.add_run()
        r.text = prefix + t
        _font(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("font", font))
    return tb


def code(s, x, y, w, h, content, size=14):
    rect(s, x, y, w, h, fill=CODE_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    text(s, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2), content,
         size=size, color=CODE_FG, font=MONO, spacing=1.0, after=1)


def chip(s, x, y, key, w=Inches(1.45), h=Inches(0.38), size=14):
    label, fg, bgc = STATUS[key]
    rect(s, x, y, w, h, fill=bgc, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=fg, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def new_slide(eyebrow, title, status=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, fill=BG)
    rect(s, 0, 0, Inches(0.14), H, fill=ACCENT)
    text(s, MX, Inches(0.32), Inches(9), Inches(0.4), eyebrow, size=14, bold=True, color=ACCENT)
    text(s, MX, Inches(0.68), Inches(10.6), Inches(0.75), title, size=28, bold=True)
    if status:
        chip(s, W - MX - Inches(1.75), Inches(0.5), status, w=Inches(1.75), h=Inches(0.45), size=15)
    s.shapes.add_connector(1, MX, Inches(1.42), W - MX, Inches(1.42)).line.color.rgb = LINE
    slides_meta.append(s)
    return s


def table(s, x, y, col_w, rows, size=15, row_h=Inches(0.46), header=True, status_col=None):
    """rows: list of lists (strings or STATUS keys when status_col)."""
    cy = y
    for ri, row in enumerate(rows):
        is_head = header and ri == 0
        cx = x
        rh = row_h
        rect(s, x, cy, sum(col_w), rh, fill=ACCENT if is_head else (CARD if ri % 2 else BG), line=None if is_head else LINE)
        for ci, cell in enumerate(row):
            if status_col is not None and ci == status_col and not is_head and cell in STATUS:
                chip(s, cx + Inches(0.08), cy + (rh - Inches(0.36)) / 2, cell, w=col_w[ci] - Inches(0.16), h=Inches(0.36), size=14)
            else:
                text(s, cx + Inches(0.08), cy, col_w[ci] - Inches(0.16), rh, cell, size=size,
                     bold=is_head, color=BG if is_head else INK, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, after=0)
            cx += col_w[ci]
        cy += rh
    return cy


def card(s, x, y, w, h, title, body, accent=ACCENT, size=15, title_size=17):
    rect(s, x, y, w, h, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    rect(s, x, y + Inches(0.18), Inches(0.07), Inches(0.42), fill=accent)
    text(s, x + Inches(0.22), y + Inches(0.14), w - Inches(0.35), Inches(0.5), title, size=title_size, bold=True)
    text(s, x + Inches(0.22), y + Inches(0.66), w - Inches(0.35), h - Inches(0.75), body, size=size, color=INK, bullets=True, after=5)



def pic(s, path, x, y, w=None, h=None, border=True):
    p = s.shapes.add_picture(str(path), x, y, w, h)
    if border:
        p.line.color.rgb = LINE
        p.line.width = Pt(1)
    return p


def caption(s, x, y, w, t):
    text(s, x, y, w, Inches(0.36), t, size=14, color=MUTED, align=PP_ALIGN.CENTER)


def callout(s, x, y, w, h, lines, fg=RED, bg=RED_LT, size=16):
    rect(s, x, y, w, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, x + Inches(0.25), y, w - Inches(0.5), h, lines, size=size, color=fg, anchor=MSO_ANCHOR.MIDDLE, after=2)


CW = W - 2 * MX  # content width
EST = "※ 시간은 모두 추정치 · '공식'만 원스토어 문서 기재값"

# =====================================================================
# 1. Title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=ACCENT)
rect(s, Inches(0.8), Inches(1.5), Inches(0.12), Inches(2.6), fill=BG)
text(s, Inches(1.15), Inches(1.4), Inches(11), Inches(0.5), "한국어 성경 · ONE store", size=20, bold=True, color=ACCENT_LT)
text(s, Inches(1.15), Inches(1.95), Inches(11.5), Inches(1.2), "원스토어 출시 프로세스", size=44, bold=True, color=BG)
text(s, Inches(1.15), Inches(3.1), Inches(11), Inches(1.0),
     "가입 → APK → 자산 → 콘솔 등록 → 심사 → 출시 · 담당 · 추정 소요시간", size=20, color=ACCENT_LT)
text(s, Inches(1.15), Inches(5.3), Inches(11), Inches(1.2),
     [f"{DATE}  ·  com.lsh.koreanbible  ·  v1.0.0 (versionCode 1)", "원본 문서: docs/project-management/ONESTORE_RELEASE_GUIDE.md"],
     size=16, color=BG)
slides_meta.append(s)

# 2. Why ONE store first
s = new_slide("WHY", "왜 원스토어를 먼저 하나")
card(s, MX, CONTENT_TOP, Inches(5.95), Inches(3.55), "Google Play (신규 개인 계정)", [
    "프로덕션 전에 비공개 테스트 필수",
    "테스터 12명 이상 × 14일 연속 opt-in",
    "그 뒤 프로덕션 액세스 신청 → 심사",
    "테스터 모집·유지가 일정의 병목",
], accent=AMBER, size=17)
card(s, MX + Inches(6.18), CONTENT_TOP, Inches(5.95), Inches(3.55), "원스토어", [
    "문서상 비공개 테스트 인원·기간 요건 없음",
    "가입 무료 · 개인개발자 가능",
    "심사 최대 5영업일 (공식)",
    "같은 서명 APK가 이미 준비됨",
], accent=GREEN, size=17)
callout(s, MX, CONTENT_TOP + Inches(3.8), CW, Inches(1.2), [
    ("Play는 그대로 병행: 12명 × 14일 테스트는 계속 진행", {"bold": True}),
    ("주의: Play 미출시 패키지는 원스토어 설치 시 Play 프로텍트 경고 가능 (10번 슬라이드)", {"size": 15}),
], fg=BLUE, bg=BLUE_LT)

# 3. Dashboard
s = new_slide("STATUS", "현재 상태 대시보드 (2026-10-03)")
kpis = [("서명 APK", "준비됨", "58.1 MB · 업로드 키 서명", GREEN),
        ("등록 자산", "초안 완료", "아이콘 · 배너 · 스샷 6장", GREEN),
        ("판매 문구", "초안 완료", "제목·설명·키워드", GREEN),
        ("개발자 계정", "미가입", "사용자 가입 필요", RED)]
kw = (CW - Inches(0.3) * 3) / 4
for i, (t, v, d, c) in enumerate(kpis):
    x = MX + i * (kw + Inches(0.3))
    rect(s, x, CONTENT_TOP + Inches(0.1), kw, Inches(1.9), fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, x + Inches(0.2), CONTENT_TOP + Inches(0.2), kw - Inches(0.4), Inches(0.4), t, size=15, color=MUTED, bold=True)
    text(s, x + Inches(0.2), CONTENT_TOP + Inches(0.6), kw - Inches(0.4), Inches(0.75), v, size=30, bold=True, color=c)
    text(s, x + Inches(0.2), CONTENT_TOP + Inches(1.35), kw - Inches(0.4), Inches(0.6), d, size=14, color=INK)
table(s, MX, CONTENT_TOP + Inches(2.3), [Inches(3.0), Inches(7.3), Inches(1.83)], [
    ["영역", "현황", "상태"],
    ["targetSdk / minSdk", "36 / 24 · 원스토어 최소 33 (2026-09-01~) 충족", "DONE"],
    ["의존성", "GMS·Play Core·Firebase·Billing·광고 없음 · 원스토어 SDK 불필요", "DONE"],
    ["개인정보", "INTERNET 권한 없음 → 수집정보 '아니오' · 방침 URL 선택", "DONE"],
    ["실기기 테스트", "Android 기기 설치 스모크 + Play 프로텍트 확인 필요", "NEXT"],
    ["가입 · 등록 · 심사", "원스토어 개발자센터에서 사용자 진행", "BLOCKED"],
], size=15, status_col=2)

# 4. Step overview
s = new_slide("OVERVIEW", "단계 한눈에 보기 · 담당 · 추정 시간")
table(s, MX, CONTENT_TOP, [Inches(0.6), Inches(4.6), Inches(2.3), Inches(2.3), Inches(2.33)], [
    ["#", "단계", "담당", "시간 (추정)", "공식 기준"],
    ["1", "개발자 회원가입 (개인개발자) + 이메일 인증", "사용자 · 브라우저", "15~30분", "—"],
    ["2", "정산정보 (무료앱 → 생략)", "—", "0분", "유료/IAP만"],
    ["3", "서명 APK 빌드 · 검증", "에이전트 · PC", "완료", "—"],
    ["4", "그래픽 자산 (아이콘 · 배너 · 스샷)", "에이전트 + 사용자", "2~4시간", "—"],
    ["5", "판매 문구", "에이전트 + 사용자", "30~60분", "—"],
    ["6", "개인정보 URL (선택)", "사용자", "10~20분", "—"],
    ["7", "ONEconsole 상품 등록", "사용자 · 브라우저", "1~1.5시간", "—"],
    ["8", "심사", "원스토어", "대기", "최대 5영업일"],
    ["9", "출시 적용", "사용자", "5분", "AAB면 ≤30분"],
    ["10", "출시 후 확인 · 업데이트", "사용자 + 에이전트", "회당 30~60분", "업데이트도 심사"],
], size=14, row_h=Inches(0.43))
text(s, MX, CONTENT_TOP + Inches(4.85), CW, Inches(0.4), EST, size=14, color=MUTED)

# 5. Timeline
s = new_slide("TIMELINE", "전체 소요: 작업 4~7시간 · 달력 1~2주 (추정)")
rect(s, MX, CONTENT_TOP + Inches(0.05), Inches(3.9), Inches(2.2), fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, MX, CONTENT_TOP + Inches(0.15), Inches(3.9), Inches(1.0), "4~7시간", size=44, bold=True, color=BG, align=PP_ALIGN.CENTER)
text(s, MX, CONTENT_TOP + Inches(1.2), Inches(3.9), Inches(0.9), "실제 작업 (추정)\n하루 안에 가능", size=16, color=ACCENT_LT, align=PP_ALIGN.CENTER)
rect(s, MX, CONTENT_TOP + Inches(2.45), Inches(3.9), Inches(2.2), fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, MX, CONTENT_TOP + Inches(2.55), Inches(3.9), Inches(1.0), "1~2주", size=44, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
text(s, MX, CONTENT_TOP + Inches(3.6), Inches(3.9), Inches(0.9), "달력 기준 (추정)\n심사 최대 5영업일 포함", size=16, color=MUTED, align=PP_ALIGN.CENTER)
# day bars
tx = MX + Inches(4.3)
tw = CW - Inches(4.3)
days = [("1일차", "가입 · 자산 확정 · 문구 · 콘솔 등록 · 심사 요청", 0.0, 0.14, BLUE),
        ("2~8일차", "심사 (공식 최대 5영업일 · 주말 제외)", 0.14, 0.72, AMBER),
        ("승인 후", "출시 적용 · 실기기 설치 확인", 0.72, 0.82, GREEN),
        ("여유", "보완 요청 · 재심사 대비", 0.82, 1.0, LINE)]
for i, (lab, desc, a, b, c) in enumerate(days):
    y = CONTENT_TOP + Inches(0.15) + i * Inches(1.2)
    text(s, tx, y, Inches(1.4), Inches(0.45), lab, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, tx + Inches(1.5), y + Inches(0.2), tw - Inches(1.5), Inches(0.05), fill=LINE)
    bw = int((tw - Inches(1.5)) * (b - a))
    rect(s, tx + Inches(1.5) + int((tw - Inches(1.5)) * a), y + Inches(0.02), max(bw, Inches(0.4)), Inches(0.42),
         fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, tx + Inches(1.5), y + Inches(0.5), tw - Inches(1.5), Inches(0.45), desc, size=14, color=INK)
text(s, MX, CONTENT_TOP + Inches(4.8), CW, Inches(0.36), "공식: 심사 최대 5영업일 (경우에 따라 더 소요) · 미국 배포 시 약 5영업일 추가 · 2024 블로그 후기 '하루 내'는 비공식", size=14, color=MUTED)

# 6. Step 1 signup
s = new_slide("STEP 1 · 사용자 · 15~30분 (추정)", "원스토어 개발자 회원가입", "BLOCKED")
steps = [
    ("1", "접속", "dev.onestore.net → 회원가입 (무료 · 만 14세 이상)"),
    ("2", "약관 · 계정", "약관 동의 → 아이디 · 비밀번호 · 보안 질문 · 국가: 대한민국"),
    ("3", "회원 유형", "개인개발자 (개인사업자 · 법인사업자 아님)"),
    ("4", "기본정보", "이름 · 휴대폰 · 이메일 · 전화 · 주소 · 생년월일"),
    ("5", "인증", "이메일 인증으로 가입 완료 (휴대폰 본인인증 화면이 나올 수 있음)"),
]
for i, (n, t, d) in enumerate(steps):
    y = CONTENT_TOP + i * Inches(0.82)
    rect(s, MX, y, Inches(0.6), Inches(0.6), fill=ACCENT, shape=MSO_SHAPE.OVAL)
    text(s, MX, y, Inches(0.6), Inches(0.6), n, size=18, bold=True, color=BG, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + Inches(0.85), y, Inches(2.0), Inches(0.6), t, size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + Inches(2.9), y, Inches(9.2), Inches(0.6), d, size=16, anchor=MSO_ANCHOR.MIDDLE)
callout(s, MX, CONTENT_TOP + Inches(4.2), CW, Inches(0.95), [
    ("정산정보(STEP 2)는 무료앱이라 생략. 유료/인앱결제 시에만 범용 공인인증서로 계좌 인증", {"bold": True}),
], fg=GREEN, bg=GREEN_LT, size=16)

# 7. Step 3 APK
s = new_slide("STEP 3 · 에이전트 · 완료", "서명된 릴리스 APK", "DONE")
table(s, MX, CONTENT_TOP, [Inches(2.6), Inches(9.53)], [
    ["항목", "값 (2026-10-03 02:08 KST 빌드)"],
    ["경로", "build\\app\\outputs\\flutter-apk\\app-release.apk"],
    ["크기", "60,871,249 bytes (≈58.1 MB) · 최대 2 GB 제한"],
    ["패키지 · 버전", "com.lsh.koreanbible · versionCode 1 · versionName 1.0.0"],
    ["SDK", "minSdk 24 · targetSdk 36"],
    ["서명", "CN=LSH, O=personal, C=KR · APK Signature Scheme v2"],
    ["SHA-256", "41:9D:AA:1C:6E:72:33:EB:FC:32:9B:7B:C4:B5:FD:FE:\n2E:74:94:C2:69:B8:7D:42:19:57:AC:AE:94:6E:33:87  = 업로드 키"],
], size=15, row_h=Inches(0.5))
code(s, MX, CONTENT_TOP + Inches(3.75), CW, Inches(1.05),
     "flutter build apk --release\n& \"$env:LOCALAPPDATA\\Android\\sdk\\build-tools\\37.0.0\\apksigner.bat\" verify --print-certs build\\app\\outputs\\flutter-apk\\app-release.apk")

# 8. Step 4 assets
s = new_slide("STEP 4 · 에이전트 초안 + 사용자 확인 · 2~4시간 (추정)", "그래픽 자산", "DRAFT")
top = CONTENT_TOP + Inches(0.05)
pic(s, ASSETS / "banner_1024x578.png", MX, top, w=Inches(4.6))
caption(s, MX, top + Inches(2.62), Inches(4.6), "배너 1024×578 (확정 · 인디고/골드)")
pic(s, ASSETS / "icon_512.png", MX, top + Inches(3.1), w=Inches(1.55))
text(s, MX + Inches(1.7), top + Inches(3.1), Inches(3.2), Inches(1.6), [
    ("아이콘 512×512", {"bold": True}),
    ("PNG 32-bit · 115 KB", {}),
    ("스샷: 2~8장 · 장당 ≤1 MB", {}),
    ("720×1280 · 6장 완료", {}),
], size=14, after=2)
shots = [("01_bible_reading.png", "성경 읽기"), ("02_version_switch_krv_kjv_asv.png", "3역본 대조"), ("06_settings_dark_mode.png", "다크 모드")]
sw = Inches(2.3)
for i, (f, cap) in enumerate(shots):
    x = MX + Inches(5.0) + i * (sw + Inches(0.12))
    pic(s, ASSETS / "screenshots" / f, x, top, w=sw)
    caption(s, x, top + Inches(4.12), sw, cap)
text(s, MX, top + Inches(4.75), CW, Inches(0.36), "폴더: docs/store-assets/onestore/ · 스샷은 웹 빌드 캡처(실기기 아님) — 필요 시 실기기 캡처로 교체", size=14, color=MUTED)

# 9. Step 5 listing text
s = new_slide("STEP 5 · 에이전트 초안 + 사용자 검토 · 30~60분 (추정)", "판매 문구", "DRAFT")
table(s, MX, CONTENT_TOP, [Inches(2.3), Inches(8.13), Inches(1.7)], [
    ["항목", "초안 (listing_ko.md)", "길이 / 제한"],
    ["앱 제목", "한국어 성경 - 개역한글·찬송가·교독문", "21 / 50"],
    ["한 줄 설명", "광고 없이 오프라인으로 읽는 개역한글 성경.\nKJV·ASV 대조, 찬송가 102곡, 교독문 51편까지 한 앱에.", "62 / 100"],
    ["상세 설명", "성경 3역본 대조 · 검색 / 찬송 102곡 · 악보 / 교독문 51 ·\n주기도문 · 사도신경 / 북마크 / 다크 · 글자 크기 · 고대비 / 수집 없음", "685 / 1,300"],
    ["키워드", "성경, 개역한글, 찬송가, 교독문, 오프라인성경, 무료성경,\nKJV, 주기도문, 사도신경, 영어성경", "10 / 1~10"],
    ["카테고리", "도서 (대안: 생활) — 콘솔 목록에서 확인", "—"],
], size=15, row_h=Inches(0.72))
callout(s, MX, CONTENT_TOP + Inches(4.5), CW, Inches(0.7), [
    ("금지: 다른 앱마켓 링크·언급 (반려 사유) · 판매자 정보(이름·이메일·전화)는 사용자 입력", {"bold": True}),
], fg=AMBER, bg=AMBER_LT, size=15)

# 10. Phone test + Play Protect
s = new_slide("등록 전 · 사용자 · 30분 (추정)", "실기기 테스트 · Play 프로텍트 위험", "NEXT")
card(s, MX, CONTENT_TOP, Inches(5.6), Inches(3.75), "실기기 스모크 테스트", [
    "APK 설치 → 첫 실행 · 강제 종료 없음",
    "비행기 모드에서 성경 · 찬송 · 교독문",
    "다크 · 큰 글씨 · 고대비 · 북마크",
    "재시작 후 설정 유지",
    "'실행 안 됨'은 대표 반려 사유",
], accent=BLUE, size=16)
card(s, MX + Inches(5.83), CONTENT_TOP, Inches(6.3), Inches(3.75), "Play 프로텍트 차단 가능성 (공식 안내)", [
    "패키지가 Google Play에 배포된 적 없거나",
    "Play와 패키지명/서명이 다르면 → 설치 차단 경고",
    "우리 앱: 아직 Play 미출시 → 해당",
    "targetSdk 36 → '안전하지 않은 앱' 경고는 해당 없음",
    "대응: Play 출시 병행 · Google 소명 · 설치 확인",
], accent=RED, size=16)
callout(s, MX, CONTENT_TOP + Inches(4.0), CW, Inches(0.9), [
    ("공식 문서의 Google 소명 링크가 비어 있음 → URL 미확인 (불명확 항목)", {"bold": True}),
], size=15)

# 11. Console registration
s = new_slide("STEP 7 · 사용자 · 1~1.5시간 (추정)", "ONEconsole 상품 등록 — 항목별 선택", "BLOCKED")
table(s, MX, CONTENT_TOP, [Inches(2.9), Inches(7.4), Inches(1.83)], [
    ["항목", "선택", "변경"],
    ["출시 유형", "정식 출시 (베타 아님)", "LOCK"],
    ["패키지명", "com.lsh.koreanbible (바이너리와 일치)", "LOCK"],
    ["외부결제 사용 여부", "사용안함", "LOCK"],
    ["바이너리 · 서명", "APK + '앱 서명 사용 안함' (AAB 전환 시 APK 복귀 불가)", "CAUTION"],
    ["광고 SDK · Android Auto", "아니오 · 아니오", "FREE"],
    ["수집정보", "데이터 수집/공유 아니오 · 지식재산권: 퍼블릭 도메인", "FREE"],
    ["연령등급", "IARC 설문 + 원스토어 앱 등급 3+ (게임위 불필요)", "FREE"],
    ["카테고리", "도서 계열 (게임↔비게임 전환만 재심사)", "FREE"],
    ["배포 국가 · 가격", "대한민국 · 무료 (미국은 +5영업일 · 방침 URL 필요)", "FREE"],
], size=14, row_h=Inches(0.47), status_col=2)
text(s, MX, CONTENT_TOP + Inches(4.85), CW, Inches(0.36), "변경 불가 = 공식 문서에 변경 불가로 명시 · 일반 = 명시 없음 · 'Google Play 패키지 네임'(선택)은 Play 출시 후 입력", size=14, color=MUTED)

# 12. Review
s = new_slide("STEP 8 · 원스토어 · 최대 5영업일 (공식)", "심사 요청 · 대표 반려 사유", "WAIT")
flow = [("심사 요청", "콘솔에서 제출"), ("시스템 검수", "바이러스 · 서명 · 버전"), ("가이드라인 심사", "기능 · 메타데이터"), ("결과 메일", "리포트 첨부")]
bw = (CW - Inches(0.25) * 3) / 4
for i, (t, d) in enumerate(flow):
    x = MX + i * (bw + Inches(0.25))
    rect(s, x, CONTENT_TOP, bw, Inches(1.25), fill=ACCENT if i == 3 else CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x, CONTENT_TOP + Inches(0.12), bw, Inches(0.5), t, size=17, bold=True, color=BG if i == 3 else INK, align=PP_ALIGN.CENTER)
    text(s, x, CONTENT_TOP + Inches(0.65), bw, Inches(0.45), d, size=14, color=BG if i == 3 else MUTED, align=PP_ALIGN.CENTER)
table(s, MX, CONTENT_TOP + Inches(1.5), [Inches(6.0), Inches(4.3), Inches(1.83)], [
    ["반려 사유 (원스토어 FAQ · 가이드라인)", "우리 앱", "상태"],
    ["실행 안 됨 · 강제 종료", "실기기 스모크 필요", "NEXT"],
    ["다른 앱마켓 링크 · 언급", "코드에 없음 · 문구에도 금지", "DONE"],
    ["인앱/외부 결제 문제", "결제 없음", "DONE"],
    ["개인정보 수집 시 동의 누락", "수집 없음", "DONE"],
    ["부적절한 메타데이터 · 테스트용 앱", "설명·스샷이 실제 기능과 일치", "PART"],
], size=15, row_h=Inches(0.5), status_col=2)

# 13. Release
s = new_slide("STEP 9 · 사용자 · 5분 (추정)", "출시 적용", "WAIT")
opts = [("즉시 적용", "승인 즉시 판매 시작"), ("직접 적용", "승인 후 내가 버튼으로 시작"), ("예약 적용", "지정한 날짜·시간에 시작")]
ow = (CW - Inches(0.5)) / 3
for i, (t, d) in enumerate(opts):
    x = MX + i * (ow + Inches(0.25))
    card(s, x, CONTENT_TOP, ow, Inches(1.7), t, [d], accent=ACCENT, size=16)
states = ["등록중", "판매대기", "판매중"]
sw2 = Inches(2.6)
for i, st in enumerate(states):
    x = MX + Inches(1.2) + i * (sw2 + Inches(0.9))
    rect(s, x, CONTENT_TOP + Inches(2.15), sw2, Inches(0.8), fill=GREEN if i == 2 else ACCENT_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, x, CONTENT_TOP + Inches(2.15), sw2, Inches(0.8), st, size=20, bold=True, color=BG if i == 2 else ACCENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        text(s, x + sw2, CONTENT_TOP + Inches(2.15), Inches(0.9), Inches(0.8), "→", size=24, bold=True, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, MX, CONTENT_TOP + Inches(3.3), CW, Inches(1.6), [
    "APK 업로드: 기기별 APK 생성 단계 없음 · AAB 업로드 시 최대 약 30분 (공식)",
    "출시 직후: 원스토어 앱에서 검색 → 실기기 설치 → Play 프로텍트 경고 여부 확인",
    "추천: 첫 출시는 '직접 적용'으로 승인 확인 후 원하는 시점에 공개",
], size=16, bullets=True, after=8)

# 14. Updates
s = new_slide("STEP 10 · 운영", "업데이트: 같은 키 · versionCode 올리기", "NEXT")
code(s, MX, CONTENT_TOP, Inches(6.5), Inches(2.6),
     "# pubspec.yaml\nversion: 1.0.1+2      # versionName+versionCode\n\nflutter build apk --release\n# apksigner로 SHA-256 41:9D…33:87 확인\n# → 콘솔에 새 바이너리 업로드 → 심사")
text(s, MX + Inches(6.8), CONTENT_TOP, Inches(5.33), Inches(3.0), [
    "업데이트는 같은 업로드 키 서명 필수",
    "versionCode는 이전보다 커야 함",
    "디버그 서명 APK는 거부됨",
    "업데이트도 심사 진행",
    "versionCode는 마켓 간 동일 운영 권장",
], size=16, bullets=True, after=8)
callout(s, MX, CONTENT_TOP + Inches(3.05), CW, Inches(1.6), [
    ("'앱 서명 사용 안함' → 키 분실 = 업데이트 불가", {"bold": True, "size": 20}),
    ("C:\\Users\\LSH\\secure\\korean-bible-upload.jks + 비밀번호 백업을 PC 밖 두 곳에 보관", {}),
], size=16)

# 15. Signature consistency (PEPK)
s = new_slide("결정 필요", "Google Play와 서명 맞추기 (PEPK)", "CAUTION")
table(s, MX, CONTENT_TOP, [Inches(3.0), Inches(4.6), Inches(4.53)], [
    ["스토어", "기기에 설치되는 서명", "결과"],
    ["원스토어 (권장안)", "우리 업로드 키 ('앱 서명 사용 안함')", "41:9D…33:87"],
    ["Play — Google 생성 키", "Google이 만든 다른 키", "서명 불일치 → 스토어 간 업데이트 불가"],
    ["Play — PEPK로 기존 키 등록", "우리 업로드 키", "서명 일치 → 스토어 간 업데이트 가능"],
], size=15, row_h=Inches(0.62))
text(s, MX, CONTENT_TOP + Inches(2.75), CW, Inches(2.2), [
    "Play Console 앱 생성 시 '앱 서명 키' 단계에서 PEPK 도구로 기존 키 업로드를 선택",
    "원스토어 AAB FAQ: 다른 마켓과 같은 앱으로 업데이트하려면 서명 키가 같아야 → 스토어 생성 키 대신 자기 키",
    "원스토어 '원스토어 생성 키' 옵션은 비추천 (Play와 영구히 다른 서명)",
    "ANDROID_RELEASE_PROCESS.md(Play) 의 'Google 생성 키 유지' 계획도 이에 맞춰 수정 필요",
], size=16, bullets=True, after=8)

# 16. Unclear
s = new_slide("OPEN QUESTIONS", "공식 문서에서 불명확한 점")
table(s, MX, CONTENT_TOP, [Inches(0.6), Inches(4.0), Inches(7.53)], [
    ["#", "항목", "내용 / 대응"],
    ["1", "카테고리 목록", "공개 문서에 없음 → 콘솔 드롭다운에서 확인"],
    ["2", "가입 본인인증", "위키 '본인인증' vs 가입 페이지 '이메일 인증' → 휴대폰 준비"],
    ["3", "방침 URL (미수집 앱)", "필수 아님으로 읽힘 → 콘솔 UI에서 확인"],
    ["4", "패키지명 분리 권장", "'분리 권장' vs Play 패키지 입력란·Play 바이너리 허용 → 동일 유지 권장"],
    ["5", "Play 프로텍트 소명", "공식 페이지의 링크가 비어 있음"],
    ["6", "평균 심사 기간", "'최대 5영업일'만 공개"],
    ["7", "배너 필수 여부", "1024×578 그래픽이 필수인지 표현 불분명 → 준비 완료"],
], size=15, row_h=Inches(0.56))

# 17. Checklist
s = new_slide("CHECKLIST", "등록 체크리스트")
groups = [
    ("준비물 (사용자)", [("개발자 계정 가입", "BLOCKED"), ("판매자 정보", "BLOCKED"), ("카테고리 확인", "WAIT"),
                      ("실기기 스모크", "NEXT"), ("방침 URL 여부", "WAIT")]),
    ("자산 (에이전트)", [("서명 APK", "DONE"), ("아이콘 512", "DONE"), ("배너 1024×578", "DONE"),
                       ("스샷 6장", "DRAFT"), ("판매 문구", "DRAFT")]),
    ("등록 직전 확인", [("SHA-256 41:9D…33:87", "DONE"), ("targetSdk ≥ 33", "DONE"), ("정식 · 외부결제 X", "LOCK"),
                      ("수집 · 광고 아니오", "WAIT"), ("IARC + 3+", "WAIT")]),
]
gw = (CW - Inches(0.5)) / 3
for gi, (gt, items) in enumerate(groups):
    x = MX + gi * (gw + Inches(0.25))
    text(s, x, CONTENT_TOP, gw, Inches(0.45), gt, size=18, bold=True, color=ACCENT)
    for ii, (label, st) in enumerate(items):
        y = CONTENT_TOP + Inches(0.6) + ii * Inches(0.72)
        rect(s, x, y, gw, Inches(0.56), fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        text(s, x + Inches(0.12), y, gw - Inches(1.6), Inches(0.56), label, size=14, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, after=0)
        chip(s, x + gw - Inches(1.45), y + Inches(0.1), st, w=Inches(1.35), h=Inches(0.36))

# 18. Next actions
s = new_slide("NEXT", "다음 행동")
card(s, MX, CONTENT_TOP, Inches(5.95), Inches(4.6), "사용자", [
    "원스토어 개발자 가입 (개인개발자)",
    "판매자 정보 확정 (이름·이메일·전화)",
    "서명 방식 · 패키지명 · 카테고리 결정",
    "실기기 설치 테스트",
    "콘솔 등록 → 심사 요청",
], accent=RED, size=18)
card(s, MX + Inches(6.18), CONTENT_TOP, Inches(5.95), Inches(4.6), "에이전트", [
    "실기기 스샷으로 교체 (원하면)",
    "미사용 자산 정리로 APK 축소 검토 (osis.xml 5.7 MB 등)",
    "Play 문서에 PEPK 서명 계획 반영",
    "업데이트 시 빌드 · 서명 검증",
], accent=BLUE, size=18)

# 19. Sources
s = new_slide("SOURCES", "출처 (2026-10-03 확인)")
table(s, MX, CONTENT_TOP, [Inches(3.6), Inches(8.53)], [
    ["내용", "URL (onestore-dev.gitbook.io/dev/…)"],
    ["가입 · 기본정보", "docs/member/sign-up · docs/member/info/basic-info"],
    ["상품 등록 · 기본정보", "docs/apps/register-app · docs/apps/product/android/main-info"],
    ["바이너리 · 서명 · Play 프로텍트", "docs/apps/product/android/binary · help/faq/apps/one-store-android-app-bundle"],
    ["판매정보 · 이미지 규격", "docs/apps/product/android/app-info · tools/icon-guide"],
    ["수집정보 · 연령등급", "docs/apps/product/common-info/main-info/data · …/common-info/age-rating"],
    ["심사 · 반려 · 출시", "docs/apps/request-for-review · help/faq/review · docs/apps/distribution-management"],
    ["targetSdk 공지", "dev.onestore.net/devpoc/support/news/noticeView.omp?noticeId=33671"],
    ["Play 비공개 테스트 12×14", "support.google.com/googleplay/android-developer/answer/14151465"],
], size=14, row_h=Inches(0.5))

# ---------- footers ----------
total = len(slides_meta)
for i, sl in enumerate(slides_meta, start=1):
    if i == 1:
        continue
    text(sl, MX, H - Inches(0.48), Inches(8), Inches(0.36), FOOTER, size=14, color=MUTED)
    text(sl, W - MX - Inches(1.5), H - Inches(0.48), Inches(1.5), Inches(0.36), f"{i} / {total}", size=14, color=MUTED, align=PP_ALIGN.RIGHT)

if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "roadmap-pptx" / "KoreanBible_ONEstore_Release_Process_2026-10-03.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out} ({total} slides)")
