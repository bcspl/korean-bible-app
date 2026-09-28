#!/usr/bin/env python3
"""Android release process deck (Korean) for 한국어 성경 — 2026-09-29.

Usage:  python docs/project-management/generate_android_release_pptx.py [output.pptx]
Default output: docs/roadmap-pptx/KoreanBible_Android_Release_Process_2026-09-29.pptx
Requires: python-pptx (pip install python-pptx)
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt, Emu

# ---------- design tokens ----------
FONT = "Malgun Gothic"
MONO = "Consolas"
BG = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LINE = RGBColor(0xE3, 0xE7, 0xEE)
CARD = RGBColor(0xF5, 0xF7, 0xFB)
ACCENT = RGBColor(0x3F, 0x51, 0xB5)      # app indigo
ACCENT_LT = RGBColor(0xE8, 0xEB, 0xF8)
CODE_BG = RGBColor(0x1E, 0x23, 0x2E)
CODE_FG = RGBColor(0xE6, 0xED, 0xF3)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
GREEN_LT = RGBColor(0xE6, 0xF4, 0xEA)
AMBER = RGBColor(0xB2, 0x6A, 0x00)
AMBER_LT = RGBColor(0xFE, 0xF3, 0xE0)
RED = RGBColor(0xC6, 0x28, 0x28)
RED_LT = RGBColor(0xFD, 0xEC, 0xEA)
BLUE = RGBColor(0x1A, 0x5F, 0xB4)
BLUE_LT = RGBColor(0xE7, 0xF0, 0xFB)

STATUS = {
    "DONE": ("완료", GREEN, GREEN_LT),
    "NEXT": ("다음", BLUE, BLUE_LT),
    "BLOCKED": ("사용자 필요", RED, RED_LT),
    "WAIT": ("대기", AMBER, AMBER_LT),
    "PART": ("진행 중", AMBER, AMBER_LT),
}

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)
CONTENT_TOP = Inches(1.55)
MIN_PT = 14

DATE = "2026-09-29"
FOOTER = "한국어 성경 · Android 출시 프로세스 · " + DATE

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
slides_meta: list = []


def _font(run, size, bold=False, color=INK, font=FONT):
    assert size >= MIN_PT, size
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
        if el is None:
            el = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1], {})
            rPr.append(el)
        el.set("typeface", font)


def rect(s, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.RECTANGLE, radius=0.06):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, content, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.15, after=4, bullets=False):
    """content: str (\n = new paragraph) or list of str / (str, dict)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    items = content if isinstance(content, list) else str(content).split("\n")
    for i, it in enumerate(items):
        t, o = (it if isinstance(it, tuple) else (it, {}))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.line_spacing = o.get("spacing", spacing)
        p.space_after = Pt(o.get("after", after))
        prefix = "• " if (bullets and not o.get("nobullet")) else ""
        r = p.add_run()
        r.text = prefix + t
        _font(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("font", font))
    return tb


def code(s, x, y, w, h, content, size=14):
    rect(s, x, y, w, h, fill=CODE_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    text(s, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2), content,
         size=size, color=CODE_FG, font=MONO, spacing=1.0, after=1)


def chip(s, x, y, key, w=Inches(1.45), h=Inches(0.38), size=14):
    label, fg, bgc = STATUS[key]
    rect(s, x, y, w, h, fill=bgc, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=fg, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def new_slide(eyebrow, title, status=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, fill=BG)
    rect(s, 0, 0, Inches(0.14), H, fill=ACCENT)
    text(s, MX, Inches(0.32), Inches(9), Inches(0.4), eyebrow, size=14, bold=True, color=ACCENT)
    text(s, MX, Inches(0.68), Inches(10.6), Inches(0.75), title, size=28, bold=True)
    if status:
        chip(s, W - MX - Inches(1.75), Inches(0.5), status, w=Inches(1.75), h=Inches(0.45), size=15)
    s.shapes.add_connector(1, MX, Inches(1.42), W - MX, Inches(1.42)).line.color.rgb = LINE
    slides_meta.append(s)
    return s


def table(s, x, y, col_w, rows, size=15, row_h=Inches(0.46), header=True, status_col=None):
    """rows: list of lists (strings or STATUS keys when status_col)."""
    cy = y
    for ri, row in enumerate(rows):
        is_head = header and ri == 0
        cx = x
        rh = row_h
        rect(s, x, cy, sum(col_w), rh, fill=ACCENT if is_head else (CARD if ri % 2 else BG), line=None if is_head else LINE)
        for ci, cell in enumerate(row):
            if status_col is not None and ci == status_col and not is_head and cell in STATUS:
                chip(s, cx + Inches(0.08), cy + (rh - Inches(0.36)) / 2, cell, w=col_w[ci] - Inches(0.16), h=Inches(0.36), size=14)
            else:
                text(s, cx + Inches(0.08), cy, col_w[ci] - Inches(0.16), rh, cell, size=size,
                     bold=is_head, color=BG if is_head else INK, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, after=0)
            cx += col_w[ci]
        cy += rh
    return cy


def card(s, x, y, w, h, title, body, accent=ACCENT, size=15, title_size=17):
    rect(s, x, y, w, h, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    rect(s, x, y + Inches(0.18), Inches(0.07), Inches(0.42), fill=accent)
    text(s, x + Inches(0.22), y + Inches(0.14), w - Inches(0.35), Inches(0.5), title, size=title_size, bold=True)
    text(s, x + Inches(0.22), y + Inches(0.66), w - Inches(0.35), h - Inches(0.75), body, size=size, color=INK, bullets=True, after=5)


# =====================================================================
# 1. Title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=ACCENT)
rect(s, Inches(0.8), Inches(1.5), Inches(0.12), Inches(2.6), fill=BG)
text(s, Inches(1.15), Inches(1.4), Inches(11), Inches(0.5), "한국어 성경 · Google Play", size=20, bold=True, color=ACCENT_LT)
text(s, Inches(1.15), Inches(1.95), Inches(11.5), Inches(1.2), "Android 출시 전체 프로세스", size=44, bold=True, color=BG)
text(s, Inches(1.15), Inches(3.1), Inches(11), Inches(1.0),
     "Phase 0–8 · 오늘까지 완료 내역 · 다음 할 일 · 11월 1.0 공개 로드맵", size=20, color=ACCENT_LT)
text(s, Inches(1.15), Inches(5.3), Inches(11), Inches(1.2),
     [f"{DATE}  ·  com.lsh.koreanbible  ·  v1.0.0+1", "원본 문서: docs/project-management/ANDROID_RELEASE_PROCESS.md"],
     size=16, color=BG)
slides_meta.append(s)

# 2. Dashboard
s = new_slide("STATUS", "현재 상태 대시보드 (2026-09-29)")
kpis = [("제품 기능", "~92%", "다역본·찬송·교독·검증 PASS", GREEN),
        ("스토어 준비", "~40%", "ID · 키스토어 · 서명 AAB 완료", AMBER),
        ("체크리스트", "R5·R7·R8", "오늘 3건 완료", GREEN),
        ("11월 공개까지", "~5주", "비공개 테스트 14일 포함", RED)]
kw = (W - 2 * MX - Inches(0.3) * 3) / 4
for i, (t, v, d, c) in enumerate(kpis):
    x = MX + i * (kw + Inches(0.3))
    rect(s, x, CONTENT_TOP + Inches(0.1), kw, Inches(1.9), fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, x + Inches(0.2), CONTENT_TOP + Inches(0.2), kw - Inches(0.4), Inches(0.4), t, size=15, color=MUTED, bold=True)
    text(s, x + Inches(0.2), CONTENT_TOP + Inches(0.6), kw - Inches(0.4), Inches(0.75), v, size=32, bold=True, color=c)
    text(s, x + Inches(0.2), CONTENT_TOP + Inches(1.35), kw - Inches(0.4), Inches(0.6), d, size=14, color=INK)
table(s, MX, CONTENT_TOP + Inches(2.35), [Inches(3.2), Inches(7.1), Inches(1.83)], [
    ["영역", "현황", "상태"],
    ["앱 ID · 서명", "com.lsh.koreanbible · 업로드 키(alias upload) · release 서명", "DONE"],
    ["빌드", "서명 AAB 56.9MB · targetSdk 36 (Play 요구 충족)", "DONE"],
    ["개인정보처리방침", "한/영 초안 + index.html 작성 · 게시(GitHub Pages) 필요", "PART"],
    ["Play Console", "개발자 계정 미생성 · 본인 인증 필요 ($25)", "BLOCKED"],
    ["리스팅 자산", "아이콘만 있음 · 피처 그래픽/스크린샷 필요", "WAIT"],
], size=15, status_col=2)

# 3. Roadmap timeline
s = new_slide("ROADMAP", "11월 1.0 공개까지 로드맵")
months = ["9월", "10월", "11월"]
tl_x, tl_w = MX + Inches(2.6), W - 2 * MX - Inches(2.6)
mw = tl_w / 3
for i, m in enumerate(months):
    rect(s, tl_x + i * mw, CONTENT_TOP, mw - Inches(0.04), Inches(0.45), fill=ACCENT_LT)
    text(s, tl_x + i * mw, CONTENT_TOP, mw, Inches(0.45), m, size=16, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
lanes = [
    ("Sprint A 기술", 0.3, 1.0, GREEN, "ID·키·AAB ✅"),
    ("Console·인증", 0.9, 1.45, RED, "계정·본인인증"),
    ("Privacy·정책", 0.9, 1.55, AMBER, "방침 URL·양식"),
    ("Sprint B UX", 1.0, 1.85, BLUE, "온보딩·역본저장·복사"),
    ("비공개 테스트", 1.45, 2.0, RED, "12명 × 14일"),
    ("자산·리스팅", 1.15, 1.95, AMBER, "아이콘·그래픽·스샷"),
    ("액세스·심사·공개", 1.95, 2.7, GREEN, "프로덕션 1.0"),
]
for i, (name, a, b, c, note) in enumerate(lanes):
    y = CONTENT_TOP + Inches(0.65) + i * Inches(0.68)
    text(s, MX, y, Inches(2.5), Inches(0.5), name, size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, tl_x, y + Inches(0.22), tl_w, Inches(0.04), fill=LINE)
    bx = tl_x + int(mw * a)
    bw = int(mw * (b - a))
    rect(s, bx, y + Inches(0.03), bw, Inches(0.44), fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    text(s, bx, y + Inches(0.03), bw, Inches(0.44), note, size=14, bold=True, color=BG, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
today_x = tl_x + int(mw * (29 / 30))
rect(s, today_x, CONTENT_TOP + Inches(0.5), Inches(0.03), Inches(4.95), fill=RED)
text(s, today_x - Inches(0.6), CONTENT_TOP + Inches(5.45), Inches(1.2), Inches(0.35), "오늘 9/29", size=14, bold=True, color=RED, align=PP_ALIGN.CENTER)

# 4. Phase overview
s = new_slide("OVERVIEW", "Phase 0–8 한눈에 보기")
phases = [
    ("0", "사전 결정", "ID ✅ · 이메일·계정·URL", "PART"),
    ("1", "Console 계정·앱 생성", "$25 · 신분증 · 기기 인증", "BLOCKED"),
    ("2", "앱 ID · 서명", "R5 · R7 완료", "DONE"),
    ("3", "품질 게이트 · 빌드", "AAB ✅ · 실기기 R11", "NEXT"),
    ("4", "리스팅 자산", "512 아이콘 · 1024×500 · 스샷", "WAIT"),
    ("5", "정책 · 개인정보", "초안 ✅ · 게시 필요", "PART"),
    ("6", "테스트 트랙", "비공개 12명 × 14일", "WAIT"),
    ("7", "액세스 · 심사 · 공개", "프로덕션 액세스 신청", "WAIT"),
    ("8", "운영", "업데이트 · Vitals · 키 백업", "NEXT"),
]
cw, ch = (W - 2 * MX - Inches(0.5)) / 3, Inches(1.62)
for i, (n, t, d, st) in enumerate(phases):
    x = MX + (i % 3) * (cw + Inches(0.25))
    y = CONTENT_TOP + (i // 3) * (ch + Inches(0.18))
    rect(s, x, y, cw, ch, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, x + Inches(0.2), y + Inches(0.22), Inches(0.62), Inches(0.62), fill=ACCENT, shape=MSO_SHAPE.OVAL)
    text(s, x + Inches(0.2), y + Inches(0.22), Inches(0.62), Inches(0.62), n, size=20, bold=True, color=BG, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.95), y + Inches(0.2), cw - Inches(1.1), Inches(0.5), t, size=17, bold=True)
    text(s, x + Inches(0.95), y + Inches(0.68), cw - Inches(1.1), Inches(0.45), d, size=14, color=MUTED)
    chip(s, x + cw - Inches(1.75), y + ch - Inches(0.52), st, w=Inches(1.55), h=Inches(0.36))

# 5. Today's work
s = new_slide("TODAY · 2026-09-28~29", "오늘 완료: R5 · R7 · R8", "DONE")
card(s, MX, CONTENT_TOP, Inches(3.9), Inches(3.3), "R5 · applicationId", [
    "com.example.korean_bible_app → com.lsh.koreanbible",
    "namespace 동일 변경",
    "MainActivity 패키지 이동",
    "커밋 b091778",
], accent=GREEN)
card(s, MX + Inches(4.1), CONTENT_TOP, Inches(3.9), Inches(3.3), "R7 · 업로드 키 + 서명", [
    "PKCS12 · RSA 2048 · alias upload",
    "유효기간 ~2054-02-14",
    "key.properties (gitignore)",
    "release signingConfig · fcdf025",
], accent=GREEN)
card(s, MX + Inches(8.2), CONTENT_TOP, Inches(3.93), Inches(3.3), "R8 · 서명 AAB", [
    "app-release.aab",
    "59,709,077 bytes (~56.9MB)",
    "UPLOAD.RSA 지문 = 업로드 키",
    "targetSdk 36 · minSdk 24",
], accent=GREEN)
text(s, MX, CONTENT_TOP + Inches(3.5), W - 2 * MX, Inches(1.4), [
    ("GitHub push 완료: ae7b2b1..fae983b  (c89d4e8 · b091778 · 6e10341 · fcdf025 · fae983b)", {"bold": True}),
    ("비밀정보 점검: key.properties·.jks 미추적, 전체 git 이력에 비밀번호 문자열 없음 확인", {}),
    ("빌드 이슈 해결: PowerShell BOM이 build.gradle.kts 컴파일을 깨뜨림 → BOM 없는 UTF-8로 저장", {"color": MUTED}),
], size=16)

# 6. Phase 0
s = new_slide("PHASE 0", "사전 결정", "PART")
table(s, MX, CONTENT_TOP, [Inches(3.4), Inches(6.9), Inches(1.83)], [
    ["결정", "값 / 메모", "상태"],
    ["Application ID (평생 고정)", "com.lsh.koreanbible", "DONE"],
    ["앱 표시명", "한국어 성경 (스토어 앱 이름 ≤ 30자)", "PART"],
    ["계정 유형", "개인(Personal) 권장 · 조직은 D-U-N-S 필요", "BLOCKED"],
    ["지원 이메일", "[지원 이메일] — 스토어·방침에 공개", "BLOCKED"],
    ["개인정보처리방침 URL", "bcspl.github.io/korean-bible-app/privacy/", "BLOCKED"],
    ["가격 · 광고", "무료 (유료 전환 불가) · 광고 없음", "DONE"],
    ["출시 경로", "내부 → 비공개(필수) → 프로덕션", "DONE"],
], size=15, row_h=Inches(0.56), status_col=2)

# 7. Phase 1 account
s = new_slide("PHASE 1", "Play Console 계정 · 본인 인증", "BLOCKED")
steps = [
    ("1", "가입", "play.google.com/console 접속 → 장기 사용할 Google 계정으로 로그인 → 계정 유형 '개인'"),
    ("2", "등록비", "US$25 1회 결제 (신용/체크카드) · 개발자 이름은 스토어에 공개됨"),
    ("3", "본인 인증", "정부 발급 신분증 제출 · 결제 프로필의 법적 이름·주소가 신분증과 정확히 일치해야 함"),
    ("4", "기기 인증", "Play Console 모바일 앱으로 실제 Android 기기(비루팅) 보유 인증"),
    ("5", "연락처 인증", "이메일·전화번호 인증 → 완료 후 앱 생성 가능"),
]
for i, (n, t, d) in enumerate(steps):
    y = CONTENT_TOP + i * Inches(0.92)
    rect(s, MX, y, Inches(0.62), Inches(0.62), fill=ACCENT, shape=MSO_SHAPE.OVAL)
    text(s, MX, y, Inches(0.62), Inches(0.62), n, size=18, bold=True, color=BG, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + Inches(0.85), y - Inches(0.02), Inches(2.0), Inches(0.62), t, size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + Inches(2.9), y - Inches(0.02), Inches(9.2), Inches(0.7), d, size=16, anchor=MSO_ANCHOR.MIDDLE)
rect(s, MX, H - Inches(1.15), W - 2 * MX, Inches(0.6), fill=RED_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
text(s, MX + Inches(0.2), H - Inches(1.15), W - 2 * MX - Inches(0.4), Inches(0.6),
     "인증 기간은 고정되어 있지 않음 → 11월 공개의 critical path. 이번 주 안에 시작 권장", size=16, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE)

# 8. Phase 1 app creation
s = new_slide("PHASE 1", "앱 만들기 (계정 승인 후)", "WAIT")
table(s, MX, CONTENT_TOP, [Inches(3.6), Inches(8.53)], [
    ["항목", "입력값"],
    ["앱 이름", "한국어 성경"],
    ["기본 언어", "한국어 – ko-KR"],
    ["앱 또는 게임", "앱"],
    ["무료 또는 유료", "무료"],
    ["선언", "개발자 프로그램 정책 · 미국 수출법 동의"],
    ["앱 서명", "Play App Signing (Google 생성 앱 서명 키) 유지"],
], size=16, row_h=Inches(0.56))
text(s, MX, CONTENT_TOP + Inches(4.15), W - 2 * MX, Inches(1.2), [
    "앱 생성 직후 대시보드의 '앱 설정' 작업 목록(정책·리스팅)을 순서대로 채움",
    "첫 업로드는 내부 테스트 트랙 → 설치 확인 후 비공개 테스트",
], size=16, bullets=True)

# 9. 12 testers
s = new_slide("PHASE 6 · 핵심 규칙", "신규 개인 계정: 비공개 테스트 12명 × 14일", "BLOCKED")
rect(s, MX, CONTENT_TOP, Inches(5.2), Inches(3.9), fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, MX, CONTENT_TOP + Inches(0.35), Inches(5.2), Inches(1.4), "12명+", size=60, bold=True, color=BG, align=PP_ALIGN.CENTER)
text(s, MX, CONTENT_TOP + Inches(1.75), Inches(5.2), Inches(1.0), "14일 연속 opt-in", size=30, bold=True, color=BG, align=PP_ALIGN.CENTER)
text(s, MX + Inches(0.3), CONTENT_TOP + Inches(2.75), Inches(4.6), Inches(1.0), "2023-11-13 이후 생성된 개인 계정", size=16, color=ACCENT_LT, align=PP_ALIGN.CENTER)
text(s, MX + Inches(5.6), CONTENT_TOP, Inches(6.5), Inches(4.2), [
    "비공개(Closed) 테스트만 인정 — 내부 테스트는 카운트 안 됨",
    "중간 opt-out 시 연속 14일이 끊김 → 15명 이상 모집 권장",
    "요건 충족 후 대시보드에서 '프로덕션 액세스 신청' (질문 답변)",
    "액세스 심사 보통 7일 이내 → 그 후 프로덕션 출시 제출",
    "테스터: 가족·교회·지인의 Gmail 주소 목록 준비",
    "11월 공개 목표 → 늦어도 10월 중순 비공개 테스트 시작",
], size=17, bullets=True, after=8)
text(s, MX, H - Inches(0.95), W - 2 * MX, Inches(0.4),
     "출처: support.google.com/googleplay/android-developer/answer/14151465 (2026-09-29 확인)", size=14, color=MUTED)

# 10. Phase 2 app id
s = new_slide("PHASE 2 · R5", "applicationId 변경", "DONE")
code(s, MX, CONTENT_TOP, Inches(6.6), Inches(2.6),
     "// android/app/build.gradle.kts\nandroid {\n    namespace = \"com.lsh.koreanbible\"\n    defaultConfig {\n        applicationId = \"com.lsh.koreanbible\"\n    }\n}")
text(s, MX + Inches(6.9), CONTENT_TOP, Inches(5.2), Inches(4.4), [
    "이전: com.example.korean_bible_app (Play 거부)",
    "MainActivity.kt → kotlin/com/lsh/koreanbible/",
    "표시명 '한국어 성경' 유지",
    "aapt 확인: package com.lsh.koreanbible",
    "iOS·macOS·Linux·Windows ID는 아직 com.example.* (Android 무관)",
    "커밋 b091778",
], size=16, bullets=True, after=8)
rect(s, MX, CONTENT_TOP + Inches(2.85), Inches(6.6), Inches(1.5), fill=AMBER_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, MX + Inches(0.2), CONTENT_TOP + Inches(2.95), Inches(6.2), Inches(1.35), [
    ("주의: PowerShell 5 BOM", {"bold": True, "color": AMBER}),
    ("Set-Content -Encoding UTF8 은 BOM을 붙여 Gradle 스크립트 컴파일 실패 → [IO.File]::WriteAllText + UTF8Encoding($false)", {"size": 15}),
], size=16)

# 11. Phase 2 keystore
s = new_slide("PHASE 2 · R7", "업로드 키스토어 생성", "DONE")
code(s, MX, CONTENT_TOP, W - 2 * MX, Inches(1.75),
     "$kt = \"C:\\Program Files\\Android\\Android Studio\\jbr\\bin\\keytool.exe\"\n"
     "& $kt -genkeypair -v -keystore C:\\Users\\LSH\\secure\\korean-bible-upload.jks `\n"
     "  -storetype PKCS12 -keyalg RSA -keysize 2048 -validity 10000 `\n"
     "  -alias upload -dname \"CN=LSH, O=personal, C=KR\" -storepass <PW> -keypass <PW>")
table(s, MX, CONTENT_TOP + Inches(1.95), [Inches(2.4), Inches(9.73)], [
    ["항목", "값"],
    ["키스토어", "C:\\Users\\LSH\\secure\\korean-bible-upload.jks (PKCS12 · RSA 2048)"],
    ["alias / 소유자", "upload  ·  CN=LSH, O=personal, C=KR  ·  ~2054-02-14"],
    ["SHA-1", "AA:71:3A:D2:54:66:1C:37:1C:3B:18:D6:60:57:65:F6:54:39:9F:A7"],
    ["SHA-256", "41:9D:AA:1C:6E:72:33:EB:FC:32:9B:7B:C4:B5:FD:FE:\n2E:74:94:C2:69:B8:7D:42:19:57:AC:AE:94:6E:33:87"],
    ["비밀번호 위치", "android\\key.properties (gitignore) · secure\\korean-bible-key.properties.backup.txt"],
], size=15, row_h=Inches(0.5))

# 12. Gradle signing + Play App Signing
s = new_slide("PHASE 2", "Gradle release 서명 · Play App Signing", "DONE")
code(s, MX, CONTENT_TOP, Inches(6.9), Inches(4.6),
     "// 요약 (전체: android/app/build.gradle.kts)\nval keystorePropertiesFile =\n    rootProject.file(\"key.properties\")\n"
     "val hasReleaseKeystore =\n    keystorePropertiesFile.exists()\n\n"
     "signingConfigs {\n  if (hasReleaseKeystore) create(\"release\") {\n"
     "    keyAlias = props[\"keyAlias\"]\n    storeFile = file(props[\"storeFile\"])\n    ...\n  }\n}\n"
     "release { signingConfig =\n  if (hasReleaseKeystore) release else debug }")
text(s, MX + Inches(7.2), CONTENT_TOP, Inches(4.9), Inches(0.45), "Play App Signing (앱 생성 후)", size=18, bold=True)
text(s, MX + Inches(7.2), CONTENT_TOP + Inches(0.55), Inches(4.9), Inches(4.2), [
    "우리 키 = 업로드 키",
    "사용자 기기용 앱 서명 키는 Google이 보관",
    "AAB 업로드 필수 · 신규 앱 기본값",
    "앱 무결성 화면에서 업로드 키 SHA-1/256 대조",
    "업로드 키 분실 시 재설정 요청 가능",
    "key.properties 없으면 debug 서명 (CI 빌드 유지)",
], size=16, bullets=True, after=8)

# 13. Key backup warning
s = new_slide("PHASE 8 · 즉시", "키스토어 백업 경고", "BLOCKED")
rect(s, MX, CONTENT_TOP, W - 2 * MX, Inches(1.35), fill=RED_LT, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, MX + Inches(0.3), CONTENT_TOP + Inches(0.1), W - 2 * MX - Inches(0.6), Inches(1.15), [
    ("키스토어와 비밀번호가 지금 이 PC 한 곳에만 있습니다.", {"bold": True, "size": 22, "color": RED}),
    ("PC 고장·분실 시 업데이트 업로드 불가 → Google 지원으로 업로드 키 재설정 필요 (대기 발생)", {"size": 16}),
], size=16)
table(s, MX, CONTENT_TOP + Inches(1.6), [Inches(2.6), Inches(9.53)], [
    ["백업 대상", "위치"],
    ["키스토어", "C:\\Users\\LSH\\secure\\korean-bible-upload.jks"],
    ["비밀번호", "C:\\Users\\LSH\\secure\\korean-bible-key.properties.backup.txt  ·  android\\key.properties"],
], size=15, row_h=Inches(0.5))
text(s, MX, CONTENT_TOP + Inches(3.3), W - 2 * MX, Inches(1.6), [
    "권장: 암호화 USB 1개 + 암호관리자(비밀번호) — 서로 다른 장소",
    "클라우드 보관 시 반드시 암호화(zip 암호/VeraCrypt 등)",
    "금지: git 커밋 · 공개 Drive 폴더 · 메일/메신저에 평문 전송",
], size=17, bullets=True, after=6)

# 14. Phase 3 quality gate
s = new_slide("PHASE 3", "버전 올리기 · 품질 게이트", "NEXT")
code(s, MX, CONTENT_TOP, Inches(5.6), Inches(1.1), "# pubspec.yaml\nversion: 1.0.0+1   # versionName+versionCode")
text(s, MX, CONTENT_TOP + Inches(1.25), Inches(5.6), Inches(3.4), [
    "업로드할 때마다 +N(versionCode) 증가",
    "예: 1.0.0+1 → +2 → 프로덕션 +3 → 1.0.1+4",
    "about/settings 화면의 하드코딩 버전 문자열도 함께 수정",
    "현재 analyze: info 19건 (오류·경고 0)",
], size=16, bullets=True, after=8)
code(s, MX + Inches(5.9), CONTENT_TOP, Inches(6.23), Inches(3.3),
     "cd \"C:\\Users\\LSH\\Bible App\\korean-bible-app\"\n"
     "flutter pub get\nflutter test\npython scripts\\verify_bible_texts.py\nflutter analyze")

# 15. Phase 3 build & device
s = new_slide("PHASE 3 · R8 / R11", "빌드 · 서명 확인 · 실기기 설치", "NEXT")
code(s, MX, CONTENT_TOP, W - 2 * MX, Inches(2.5),
     "flutter build appbundle --release   # → build\\app\\outputs\\bundle\\release\\app-release.aab\n"
     "flutter build apk --release         # → build\\app\\outputs\\flutter-apk\\app-release.apk\n"
     "& \"$jbr\\jarsigner.exe\" -verify -verbose -certs build\\app\\outputs\\bundle\\release\\app-release.aab\n"
     "$adb = \"$env:LOCALAPPDATA\\Android\\sdk\\platform-tools\\adb.exe\"\n"
     "& $adb install -r build\\app\\outputs\\flutter-apk\\app-release.apk")
chip(s, MX, CONTENT_TOP + Inches(2.75), "DONE")
text(s, MX + Inches(1.6), CONTENT_TOP + Inches(2.72), Inches(10.5), Inches(0.45), "R8 서명 AAB 56.9MB · 지문 일치", size=16, anchor=MSO_ANCHOR.MIDDLE)
chip(s, MX, CONTENT_TOP + Inches(3.3), "NEXT")
text(s, MX + Inches(1.6), CONTENT_TOP + Inches(3.27), Inches(10.5), Inches(0.45),
     "R11 실기기: 비행기모드 · 다크 · 큰 글씨 · 고대비 · 성경/찬송/교독/북마크 · 재시작 후 설정 유지", size=16, anchor=MSO_ANCHOR.MIDDLE)
text(s, MX, CONTENT_TOP + Inches(3.95), W - 2 * MX, Inches(0.6),
     "참고: 기존 com.example 설치본은 다른 앱으로 인식 → 나란히 설치됨, 수동 삭제", size=15, color=MUTED)

# 16. Phase 4 assets
s = new_slide("PHASE 4", "스토어 리스팅 자산 규격", "WAIT")
table(s, MX, CONTENT_TOP, [Inches(2.4), Inches(7.9), Inches(1.83)], [
    ["자산", "요구 사항", "상태"],
    ["앱 아이콘", "512×512 · 32-bit PNG(알파) · ≤ 1MB", "PART"],
    ["피처 그래픽", "1024×500 · JPEG 또는 24-bit PNG(알파 없음) · 필수", "WAIT"],
    ["휴대전화 스샷", "2–8장 · 변 320–3840px · 긴 변 ≤ 짧은 변×2 · 권장 1080×1920", "WAIT"],
    ["태블릿 스샷", "선택 (태블릿 지원 시 권장)", "WAIT"],
    ["앱 이름", "≤ 30자 · 한국어 성경", "PART"],
    ["설명", "간단 ≤ 80자 · 자세히 ≤ 4000자", "WAIT"],
    ["카테고리", "도서/참고자료 · 연락처 이메일 필수", "BLOCKED"],
], size=15, row_h=Inches(0.52), status_col=2)
text(s, MX, H - Inches(0.95), W - 2 * MX, Inches(0.4),
     "출처: support.google.com/googleplay/android-developer/answer/9866151", size=14, color=MUTED)

# 17. Phase 4 copy
s = new_slide("PHASE 4", "리스팅 카피 · 스크린샷 구성", "WAIT")
card(s, MX, CONTENT_TOP, Inches(6.0), Inches(4.5), "카피 초안", [
    "간단: 완전 오프라인 한국어 성경 · 찬송 · 예배자료 (광고·계정 없음)",
    "자세히: KRV·KJV·ASV 병렬 (최대 3역본)",
    "공개 도메인 찬송 102곡 · 교독문 51편 · 사도신경/주기도문",
    "개역개정 미수록 명시 · 한국찬송가 공식 아님",
], size=17)
card(s, MX + Inches(6.25), CONTENT_TOP, Inches(5.88), Inches(4.5), "스크린샷 (4–8장)", [
    "① 성경 본문 (라이트)",
    "② 3역본 병렬 보기",
    "③ 찬송 목록/가사",
    "④ 교독문",
    "⑤ 다크 모드 · 큰 글씨",
    "adb shell screencap -p /sdcard/s1.png",
], size=17)

# 18. Privacy findings
s = new_slide("PHASE 5 · R2", "개인정보처리방침 — 코드 조사 결과", "PART")
table(s, MX, CONTENT_TOP, [Inches(3.3), Inches(8.83)], [
    ["확인 항목", "결과 (2026-09-29)"],
    ["네트워크", "release 매니페스트에 INTERNET 권한 없음 (debug/profile만)"],
    ["제3자 SDK", "광고·분석·크래시 SDK 없음"],
    ["권한", "위험 권한 없음 · 계정/로그인 없음"],
    ["SharedPreferences", "다크모드 · 글꼴 배율 · 고대비 (기기 내)"],
    ["Hive (기기 내)", "북마크(책/장/절·미리보기·시각) · 찬송 즐겨찾기 · 본문 캐시"],
    ["자동 백업", "allowBackup 미설정 → 사용자 Google 백업 포함 가능 (방침에 명시)"],
], size=15, row_h=Inches(0.55))
text(s, MX, CONTENT_TOP + Inches(4.05), W - 2 * MX, Inches(1.0), [
    "파일: docs/privacy/privacy-policy.md (한/영) · docs/privacy/index.html · 호스팅 안내 docs/privacy/README.md",
    "게시 예정 URL: https://bcspl.github.io/korean-bible-app/privacy/  (Settings → Pages → main /docs)",
], size=15, bullets=True)

# 19. Data safety
s = new_slide("PHASE 5", "Data safety 양식 · 앱 내 방침", "NEXT")
table(s, MX, CONTENT_TOP, [Inches(5.6), Inches(6.53)], [
    ["질문", "예상 답변"],
    ["사용자 데이터 수집 또는 공유?", "아니요 (기기 밖 전송 없음)"],
    ["전송 중 암호화", "해당 없음"],
    ["계정 생성 / 삭제 URL", "계정 없음 → 불필요"],
    ["데이터 삭제 방법", "앱 삭제 또는 앱 데이터 삭제"],
    ["개인정보처리방침 링크", "필수 (미수집 앱도)"],
], size=16, row_h=Inches(0.55))
rect(s, MX, CONTENT_TOP + Inches(3.55), W - 2 * MX, Inches(1.2), fill=BLUE_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, MX + Inches(0.25), CONTENT_TOP + Inches(3.62), W - 2 * MX - Inches(0.5), Inches(1.1), [
    ("Play 사용자 데이터 정책: Console 링크 + 앱 내부에도 방침 링크 또는 텍스트", {"bold": True, "color": BLUE}),
    ("→ 설정/정보 화면에 '개인정보처리방침' 텍스트 항목 추가 (B4 라이선스 화면과 함께)", {}),
], size=16)

# 20. Other declarations + target API
s = new_slide("PHASE 5", "앱 콘텐츠 선언 · 타겟 API", "NEXT")
table(s, MX, CONTENT_TOP, [Inches(3.3), Inches(4.2)], [
    ["선언", "답변"],
    ["광고", "없음"],
    ["앱 액세스", "로그인 없이 전체 이용"],
    ["콘텐츠 등급", "IARC 설문 → 전체이용가 예상"],
    ["타겟 연령", "13세 이상 권장"],
    ["뉴스·정부·금융·건강", "해당 없음"],
], size=15, row_h=Inches(0.58))
rect(s, MX + Inches(7.8), CONTENT_TOP, Inches(4.33), Inches(3.5), fill=GREEN_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, MX + Inches(8.0), CONTENT_TOP + Inches(0.15), Inches(3.95), Inches(3.3), [
    ("타겟 API 요구", {"bold": True, "size": 18, "color": GREEN}),
    ("2026-08-31부터 신규 앱·업데이트는 API 36 (Android 16) 이상", {}),
    ("연장 요청 시 2026-11-01까지", {}),
    ("현재 targetSdk 36 ✅", {"bold": True}),
], size=16, after=8)
text(s, MX, CONTENT_TOP + Inches(3.85), W - 2 * MX, Inches(0.9), [
    "13세 미만 포함 시 가족(Families) 정책 추가 요건 적용 — 전 연령 이용 자체는 가능",
    "출처: support.google.com/googleplay/android-developer/answer/11926878",
], size=14, color=MUTED)

# 21. Testing tracks
s = new_slide("PHASE 6", "테스트 트랙 운영 절차", "WAIT")
table(s, MX, CONTENT_TOP, [Inches(2.4), Inches(3.4), Inches(2.4), Inches(3.93)], [
    ["트랙", "목적", "인원", "비고"],
    ["내부", "업로드·설치 확인", "≤ 100", "요건에 미포함"],
    ["비공개", "신규 개인 계정 필수", "≥ 12 · 14일", "15명+ 권장"],
    ["공개", "선택", "제한 없음", ""],
    ["프로덕션", "일반 공개", "", "액세스 승인 후"],
], size=15, row_h=Inches(0.5))
text(s, MX, CONTENT_TOP + Inches(2.75), W - 2 * MX, Inches(2.3), [
    "① 내부 테스트 → 새 버전 → app-release.aab 업로드 → 출시 노트(ko-KR) → 출시",
    "② 비공개 테스트 트랙 생성 → 국가: 대한민국 → 테스터 이메일 목록 → 출시",
    "③ 테스터에게 참여 링크 → '테스터 되기' → 14일 유지 + 실제 사용·피드백",
    "④ 수정본은 versionCode 올려 재업로드",
], size=16, after=8)

# 22. Phase 7
s = new_slide("PHASE 7", "프로덕션 액세스 · 심사 · 공개", "WAIT")
flow = [("요건 충족", "12명 × 14일"), ("액세스 신청", "테스트·피드백·준비 답변"), ("액세스 승인", "보통 ≤ 7일"),
        ("프로덕션 제출", "AAB · 국가 · 노트"), ("앱 심사", "수일 (신규는 더 길 수 있음)"), ("1.0 공개", "11월 목표")]
bw = (W - 2 * MX - Inches(0.25) * 5) / 6
for i, (t, d) in enumerate(flow):
    x = MX + i * (bw + Inches(0.25))
    rect(s, x, CONTENT_TOP + Inches(0.3), bw, Inches(2.0), fill=ACCENT if i == 5 else CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x + Inches(0.1), CONTENT_TOP + Inches(0.45), bw - Inches(0.2), Inches(0.8), t, size=17, bold=True,
         color=BG if i == 5 else INK, align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.1), CONTENT_TOP + Inches(1.2), bw - Inches(0.2), Inches(1.0), d, size=14,
         color=BG if i == 5 else MUTED, align=PP_ALIGN.CENTER)
text(s, MX, CONTENT_TOP + Inches(2.7), W - 2 * MX, Inches(2.0), [
    "액세스 신청 답변은 구체적으로: 테스터 모집 방법, 받은 피드백, 반영한 개선, 출시 준비 상태",
    "반려 시: 사유(정책/메타데이터/권한) 확인 → 수정 → versionCode 올려 재제출",
    "게시 전 대시보드의 '앱 설정' 필수 항목이 모두 완료 표시인지 확인",
], size=16, bullets=True, after=8)

# 23. Phase 8 ops
s = new_slide("PHASE 8", "운영 · 업데이트", "NEXT")
code(s, MX, CONTENT_TOP, Inches(6.4), Inches(2.6),
     "# 1) pubspec.yaml: version 1.0.1+5\n# 2) 품질 게이트\nflutter test; flutter analyze\npython scripts\\verify_bible_texts.py\n# 3) 빌드\nflutter build appbundle --release")
text(s, MX + Inches(6.7), CONTENT_TOP, Inches(5.4), Inches(4.5), [
    "Console → 새 버전 → AAB → 출시 노트",
    "단계적 출시 (예: 20% → 100%)",
    "Android Vitals: 비정상 종료 · ANR · 시작 시간",
    "리뷰 응답 · 평점 모니터링",
    "매년 8월 말 타겟 API 상향 확인",
    "키스토어 백업 상태 분기별 점검",
], size=16, bullets=True, after=8)

# 24. Backward schedule
s = new_slide("SCHEDULE", "11월 공개를 위한 역산 일정", "NEXT")
table(s, MX, CONTENT_TOP, [Inches(2.2), Inches(8.1), Inches(1.83)], [
    ["목표일", "할 일", "상태"],
    ["~10/05", "Console 가입 · 본인 인증 시작 · 이메일 확정 · Privacy URL 게시", "BLOCKED"],
    ["~10/10", "앱 생성 · 앱 콘텐츠 작성 · 내부 테스트 AAB 업로드", "NEXT"],
    ["~10/14", "비공개 테스트 시작 (12명+ opt-in) ← 가장 중요", "BLOCKED"],
    ["10월", "Sprint B UX (온보딩·역본저장·복사·라이선스/방침 화면) · 스샷", "NEXT"],
    ["~10/28", "14일 충족 → 프로덕션 액세스 신청", "WAIT"],
    ["~11/04", "승인 → 프로덕션 제출", "WAIT"],
    ["11월 중", "심사 통과 · 1.0 공개", "WAIT"],
], size=15, row_h=Inches(0.53), status_col=2)

# 25. Checklist summary
s = new_slide("CHECKLIST", "R1–R25 요약")
groups = [
    ("정책·고지", [("R1 스토어 설명", "WAIT"), ("R2 개인정보 URL", "PART"), ("R3 콘텐츠 등급", "WAIT"), ("R4 찬송 고지", "DONE")]),
    ("Android 기술", [("R5 applicationId", "DONE"), ("R6 표시명", "PART"), ("R7 키스토어", "DONE"), ("R8 AAB", "DONE"),
                      ("R9 targetSdk 36", "DONE"), ("R10 64-bit", "WAIT"), ("R11 실기기", "NEXT")]),
    ("품질·자산·UX", [("R12 test · R13 검증", "DONE"), ("R14–17 스모크", "WAIT"), ("R18–20 자산·카피", "WAIT"),
                     ("R21 이메일·계정", "BLOCKED"), ("R22–25 사전 UX", "WAIT")]),
]
gw = (W - 2 * MX - Inches(0.5)) / 3
for gi, (gt, items) in enumerate(groups):
    x = MX + gi * (gw + Inches(0.25))
    text(s, x, CONTENT_TOP, gw, Inches(0.45), gt, size=18, bold=True, color=ACCENT)
    for ii, (label, st) in enumerate(items):
        y = CONTENT_TOP + Inches(0.6) + ii * Inches(0.64)
        rect(s, x, y, gw, Inches(0.52), fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        text(s, x + Inches(0.12), y, gw - Inches(1.55), Inches(0.52), label, size=14, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, after=0)
        chip(s, x + gw - Inches(1.4), y + Inches(0.08), st, w=Inches(1.3), h=Inches(0.36))

# 26. Next actions
s = new_slide("NEXT", "다음 행동")
card(s, MX, CONTENT_TOP, Inches(5.95), Inches(4.7), "사용자 (차단 해제)", [
    "Play Console 가입 · $25 · 신분증 · 기기 인증",
    "지원 이메일 확정",
    "GitHub Pages 켜기 (main /docs) → 방침 URL 확인",
    "키스토어 + 비밀번호 PC 밖 백업",
    "비공개 테스터 15명 Gmail 모집",
], accent=RED, size=18)
card(s, MX + Inches(6.18), CONTENT_TOP, Inches(5.95), Inches(4.7), "에이전트 (바로 가능)", [
    "R11 실기기 설치 스모크 (adb 연결 시)",
    "앱 내 개인정보처리방침·라이선스 화면",
    "Sprint B: 온보딩 · 역본 저장 · 절 복사",
    "512 아이콘 · 1024×500 피처 그래픽 · 스샷",
    "스토어 설명 최종 카피 초안",
], accent=BLUE, size=18)

# 27. Sources
s = new_slide("SOURCES", "확인한 정책 출처 (2026-09-29)")
table(s, MX, CONTENT_TOP, [Inches(4.4), Inches(7.73)], [
    ["내용", "URL (support.google.com/googleplay/android-developer/…)"],
    ["비공개 테스트 12명 × 14일", "answer/14151465"],
    ["등록비 $25 · 본인/기기 인증", "answer/6112435"],
    ["타겟 API 36 (8/31~, 연장 11/1)", "answer/11926878"],
    ["Data safety (미수집 앱도 작성)", "answer/10787469"],
    ["방침: Console + 앱 내 링크/텍스트", "answer/10144311"],
    ["그래픽 자산 규격", "answer/9866151"],
], size=15, row_h=Inches(0.55))
text(s, MX, CONTENT_TOP + Inches(4.1), W - 2 * MX, Inches(0.8),
     "정책은 수시로 변경 — 각 Phase 착수 시 재확인. 상세: docs/project-management/ANDROID_RELEASE_PROCESS.md", size=15, color=MUTED)

# ---------- footers ----------
total = len(slides_meta)
for i, sl in enumerate(slides_meta, start=1):
    if i == 1:
        continue
    text(sl, MX, H - Inches(0.48), Inches(8), Inches(0.36), FOOTER, size=14, color=MUTED)
    text(sl, W - MX - Inches(1.5), H - Inches(0.48), Inches(1.5), Inches(0.36), f"{i} / {total}", size=14, color=MUTED, align=PP_ALIGN.RIGHT)

if __name__ == "__main__":
    root = Path(__file__).resolve().parents[2]
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "docs" / "roadmap-pptx" / "KoreanBible_Android_Release_Process_2026-09-29.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out} ({total} slides)")
